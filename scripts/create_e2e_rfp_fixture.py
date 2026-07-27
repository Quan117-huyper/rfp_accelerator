"""Create a realistic RFP fixture for the full Foundry proposal workflow."""

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "generated" / "test-fixtures" / "contoso_proposal_automation_rfp.docx"


def add_bullets(document: Document, items: list[str]) -> None:
    for item in items:
        document.add_paragraph(item, style="List Bullet")


def add_heading(document: Document, text: str) -> None:
    document.add_heading(text, level=1)


def build_document() -> Path:
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

    styles = document.styles
    styles["Normal"].font.name = "Aptos"
    styles["Normal"].font.size = Pt(10)
    styles["Heading 1"].font.name = "Aptos Display"
    styles["Heading 1"].font.size = Pt(15)
    styles["Heading 1"].font.color.rgb = RGBColor(0x00, 0x58, 0x8A)

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("Request for Proposal\nProposal Automation Platform POC")
    run.bold = True
    run.font.name = "Aptos Display"
    run.font.size = Pt(22)
    run.font.color.rgb = RGBColor(0x00, 0x58, 0x8A)
    subtitle = document.add_paragraph("Customer: Contoso Manufacturing | Classification: Confidential")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER

    table = document.add_table(rows=4, cols=2)
    table.style = "Table Grid"
    for row, values in zip(table.rows, [
        ("Document version", "1.0"),
        ("Issued", "2026-07-27"),
        ("Target delivery", "POC acceptance within 8 weeks"),
        ("Primary cloud", "Microsoft Azure"),
    ]):
        row.cells[0].text, row.cells[1].text = values

    add_heading(document, "1. Business Context and Objectives")
    document.add_paragraph(
        "Contoso Manufacturing prepares customer proposals by manually reading RFP documents, "
        "collecting supporting material from SharePoint and Salesforce, and assembling PowerPoint decks."
    )
    add_bullets(document, [
        "Reduce proposal preparation time by 40 percent within six months.",
        "Improve proposal quality and consistency across proposal teams.",
        "Create an auditable approval path for requirement extraction, architecture decisions, and proposal export.",
    ])

    add_heading(document, "2. Scope and Functional Requirements")
    add_bullets(document, [
        "Proposal managers shall upload customer RFP documents in PDF, DOCX, TXT, or Markdown formats.",
        "The platform shall extract business goals, functional requirements, non-functional requirements, constraints, assumptions, integrations, security needs, and expected users.",
        "Each extracted requirement shall retain source traceability to the document section and page where available.",
        "A solution architect shall review and approve extracted requirements before architecture generation starts.",
        "The platform shall generate customer-ready proposal content and export it into approved enterprise PowerPoint templates.",
        "The platform shall support FPT FAP, Malaysia, and Singapore templates through a template mapping layer.",
    ])

    add_heading(document, "3. Document Processing")
    document.add_paragraph(
        "Digitally generated PDFs should be processed using native text parsing where possible. "
        "Scanned PDFs, complex forms, and documents containing material tables require OCR and layout extraction."
    )
    processing = document.add_table(rows=1, cols=3)
    processing.style = "Table Grid"
    processing.rows[0].cells[0].text = "Document type"
    processing.rows[0].cells[1].text = "Expected handling"
    processing.rows[0].cells[2].text = "Acceptance target"
    for values in [
        ("DOCX", "Paragraph and table extraction", "Tables retained as structured rows"),
        ("Digital PDF", "Native text extraction", "No OCR unless text is unavailable"),
        ("Scanned PDF", "OCR and layout extraction", "Readable text and tables returned"),
    ]:
        cells = processing.add_row().cells
        for index, value in enumerate(values):
            cells[index].text = value

    add_heading(document, "4. Integrations")
    add_bullets(document, [
        "The platform shall read approved reference documents from Microsoft SharePoint through Microsoft Graph.",
        "Generated proposal files shall be uploaded back to the relevant Microsoft SharePoint project folder.",
        "The platform shall read customer account context from Salesforce through a read-only API integration.",
        "The platform shall read approved reference documents from Microsoft SharePoint through Microsoft Graph.",
    ])

    add_heading(document, "5. Security and Access Control")
    add_bullets(document, [
        "Customer documents are confidential and may only be accessed by authorized project members.",
        "Employees must authenticate using enterprise single sign-on.",
        "The platform must enforce project-scoped role-based access control for proposal managers, solution architects, and administrators.",
        "The platform shall preserve audit logs for document access, requirement approval, architecture approval, and PowerPoint export.",
        "Customer document content must not be sent to public web search tools.",
        "The platform must enforce project-scoped role-based access control for proposal managers, solution architects, and administrators.",
    ])

    add_heading(document, "6. Performance, Scale, and Retention")
    add_bullets(document, [
        "The POC must process a 100-page requirement document within 10 minutes.",
        "The initial deployment shall support 30 proposal managers, 5 solution architects, and 2 administrators.",
        "Expected peak concurrency is 10 users.",
        "Uploaded customer source documents must be deleted after 30 days.",
        "Uploaded customer source documents must be retained for at least 90 days.",
    ])

    add_heading(document, "7. Constraints and Assumptions")
    add_bullets(document, [
        "Microsoft Azure is the required primary cloud platform for this POC.",
        "Monthly non-model operating cost should remain below 250 USD; Azure OpenAI token usage must be reported separately.",
        "It is assumed that the customer will provide read-only SharePoint access before UAT.",
        "The POC should use a single primary deployment region; availability targets are to be confirmed during solution design.",
    ])

    add_heading(document, "8. Out of Scope")
    add_bullets(document, [
        "Mobile applications, voice commands, gamification, blockchain approvals, and proposal trend dashboards are outside the current POC scope.",
        "The solution is an internal enterprise platform and is not a public-facing customer portal.",
    ])

    add_heading(document, "Appendix A. Repeated Source Statements")
    document.add_paragraph(
        "Architecture generation must not begin until a solution architect has approved the extracted requirements."
    )
    document.add_paragraph(
        "Microsoft SharePoint is the approved source for customer project reference documents."
    )
    document.add_paragraph(
        "The solution must support enterprise PowerPoint templates without breaking approved branding."
    )

    document.save(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    print(build_document())
