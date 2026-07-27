from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "templates" / "master_template" / "template.pptx"

BG = RGBColor(255, 255, 255)
PANEL = RGBColor(245, 247, 250)
BORDER = RGBColor(225, 230, 239)
TEXT = RGBColor(31, 41, 55)
MUTED = RGBColor(102, 102, 102)
ACCENT = RGBColor(31, 78, 121)
ACCENT_2 = RGBColor(46, 117, 182)
TEAL = RGBColor(0, 166, 166)
GREEN = RGBColor(46, 139, 87)


def add_box(slide, left, top, width, height, text="", fill=PANEL, line=BORDER, font_size=16, shape_type=MSO_AUTO_SHAPE_TYPE.RECTANGLE):
    shape = slide.shapes.add_shape(
        shape_type,
        Inches(left),
        Inches(top),
        Inches(width),
        Inches(height),
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = line
    shape.text = text
    for paragraph in shape.text_frame.paragraphs:
        for run in paragraph.runs:
            run.font.name = "Aptos"
            run.font.size = Pt(font_size)
            run.font.color.rgb = TEXT
    return shape


def add_text(slide, left, top, width, height, text="", font_size=16, color=TEXT):
    shape = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    shape.text = text
    for paragraph in shape.text_frame.paragraphs:
        for run in paragraph.runs:
            run.font.name = "Aptos"
            run.font.size = Pt(font_size)
            run.font.color.rgb = color
    return shape


def set_bg(slide):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = BG


def ensure_shapes(slide, count):
    while len(slide.shapes) < count:
        add_text(slide, 0.2, 0.2, 0.1, 0.1, "")


def build_template():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]

    # 1 cover, indexes 2-5 are filled by renderer. Matches the legacy master:
    # white canvas, large blue/teal circles, and left-aligned title stack.
    slide = prs.slides.add_slide(blank)
    set_bg(slide)
    add_box(slide, 10.6, 4.6, 3.6, 3.6, "", fill=ACCENT_2, line=ACCENT_2, shape_type=MSO_AUTO_SHAPE_TYPE.OVAL)
    add_box(slide, 11.6, 5.6, 2.0, 2.0, "", fill=TEAL, line=TEAL, shape_type=MSO_AUTO_SHAPE_TYPE.OVAL)
    add_text(slide, 0.6, 2.5, 9.5, 0.4, "ENTERPRISE PROPOSAL TEMPLATE SYSTEM", 16, ACCENT)
    add_text(slide, 0.6, 2.95, 9.5, 1.1, "Deck title", 36, TEXT)
    add_text(slide, 0.6, 4.22, 9.5, 0.55, "Subtitle", 18, MUTED)
    add_text(slide, 0.6, 6.6, 9.5, 0.4, "Prepared by", 14, MUTED)
    ensure_shapes(slide, 6)

    # 2 section opener.
    slide = prs.slides.add_slide(blank)
    set_bg(slide)
    add_box(slide, 0.6, 0.7, 0.55, 0.55, "", fill=TEAL, line=TEAL, shape_type=MSO_AUTO_SHAPE_TYPE.OVAL)
    add_text(slide, 0.6, 0.7, 0.55, 0.55, "01", 14, RGBColor(255, 255, 255))
    add_text(slide, 1.35, 0.78, 3.0, 0.4, "SECTION", 16, ACCENT)
    add_text(slide, 0.6, 3.1, 12.13, 1.2, "Solution Overview", 40, TEXT)
    add_text(slide, 0.6, 4.15, 9.5, 0.6, "Subtitle", 20, MUTED)
    ensure_shapes(slide, 5)

    # 3 agenda, renderer fills 3,6,9,12,15,18,21.
    slide = prs.slides.add_slide(blank)
    set_bg(slide)
    add_text(slide, 0.6, 0.45, 12.13, 0.6, "Agenda", 30, TEXT)
    for index in range(7):
        y = 1.75 + index * 0.66
        color = ACCENT if index % 2 == 0 else ACCENT_2
        add_box(slide, 0.6, y, 0.42, 0.42, "", fill=color, line=color, shape_type=MSO_AUTO_SHAPE_TYPE.OVAL)
        add_text(slide, 0.6, y, 0.42, 0.42, f"{index + 1:02d}", 10, RGBColor(255, 255, 255))
        add_text(slide, 1.25, y - 0.04, 9.5, 0.5, f"Agenda {index + 1}", 20, TEXT)
    ensure_shapes(slide, 22)

    # 4 executive cards.
    slide = prs.slides.add_slide(blank)
    set_bg(slide)
    add_text(slide, 0.7, 0.55, 5.0, 0.4, "Executive Summary", 24, ACCENT)
    for index in range(3):
        x = 0.8 + index * 4.15
        add_box(slide, x, 1.55, 3.65, 4.4, "", fill=PANEL)
        badge_color = [ACCENT, ACCENT_2, GREEN][index]
        add_box(slide, x + 0.18, 1.78, 0.5, 0.5, str(index + 1), fill=badge_color, line=badge_color, shape_type=MSO_AUTO_SHAPE_TYPE.OVAL)
        add_text(slide, x + 0.25, 2.55, 3.0, 0.45, "Card title", 18)
        add_text(slide, x + 0.25, 3.1, 3.0, 1.6, "Card body", 14, MUTED)
    ensure_shapes(slide, 16)

    # 5 why this solution.
    slide = prs.slides.add_slide(blank)
    set_bg(slide)
    add_text(slide, 0.7, 0.6, 8.0, 0.45, "Why this solution", 22, ACCENT)
    add_text(slide, 0.7, 1.18, 9.8, 0.45, "Subtitle", 18)
    add_box(slide, 0.8, 1.95, 11.7, 4.6, "Bullets", fill=PANEL, font_size=18)
    ensure_shapes(slide, 3)

    # 6 business / technical split.
    slide = prs.slides.add_slide(blank)
    set_bg(slide)
    add_text(slide, 0.7, 0.55, 5.0, 0.4, "Solution Split", 24, ACCENT)
    add_box(slide, 0.8, 1.55, 5.75, 4.9, "", fill=PANEL)
    add_text(slide, 1.05, 1.9, 5.1, 0.4, "Business", 20)
    add_text(slide, 1.05, 2.45, 5.1, 2.8, "Business bullets", 16, MUTED)
    add_text(slide, 7.0, 1.9, 5.1, 0.4, "Technical", 20)
    add_text(slide, 7.0, 2.45, 5.1, 2.8, "Technical bullets", 16, MUTED)
    ensure_shapes(slide, 6)

    # 7 snapshot cards.
    slide = prs.slides.add_slide(blank)
    set_bg(slide)
    add_text(slide, 0.7, 0.55, 5.0, 0.4, "Engagement Snapshot", 24, ACCENT)
    for index in range(4):
        x = 0.8 + (index % 2) * 6.0
        y = 1.35 + (index // 2) * 2.25
        add_box(slide, x, y, 5.3, 1.65, "", fill=PANEL)
        add_text(slide, x + 0.25, y + 0.25, 4.8, 0.35, "Key", 18)
        add_text(slide, x + 0.25, y + 0.72, 4.8, 0.55, "Value", 15, MUTED)
    ensure_shapes(slide, 17)

    # 8 KPIs.
    slide = prs.slides.add_slide(blank)
    set_bg(slide)
    add_text(slide, 0.7, 0.55, 5.0, 0.4, "Target Outcomes", 24, ACCENT)
    for index in range(3):
        x = 1.0 + index * 4.0
        add_box(slide, x, 1.7, 3.15, 3.8, "", fill=PANEL)
        add_text(slide, x + 0.25, 2.0, 2.6, 0.65, "Metric", 28, ACCENT)
        add_text(slide, x + 0.25, 2.85, 2.6, 0.45, "Status", 18)
        add_text(slide, x + 0.25, 3.45, 2.6, 0.9, "Label", 15, MUTED)
    ensure_shapes(slide, 13)

    # 9 workflow.
    slide = prs.slides.add_slide(blank)
    set_bg(slide)
    add_text(slide, 0.7, 0.55, 5.0, 0.4, "Workflow", 24, ACCENT)
    for index in range(5):
        x = 0.6 + index * 2.55
        add_box(slide, x, 1.65, 2.15, 3.9, "", fill=PANEL)
        add_text(slide, x + 0.2, 2.0, 1.75, 0.35, "Step", 17)
        add_text(slide, x + 0.2, 2.55, 1.75, 1.35, "Description", 13, MUTED)
    ensure_shapes(slide, 22)

    # 10 timeline. Renderer reads label at title_index - 1 and writes same text index.
    slide = prs.slides.add_slide(blank)
    set_bg(slide)
    add_text(slide, 0.7, 0.55, 5.0, 0.4, "Timeline", 24, ACCENT)
    for index in range(4):
        x = 0.9 + index * 3.05
        add_text(slide, x, 1.45, 1.0, 0.3, f"W{index + 1}", 16, ACCENT)
        add_box(slide, x, 1.95, 2.4, 3.0, "Milestone", fill=PANEL, font_size=15)
    ensure_shapes(slide, 14)

    # 11 costs. Shape 1 must be table.
    slide = prs.slides.add_slide(blank)
    set_bg(slide)
    add_text(slide, 0.7, 0.55, 5.0, 0.4, "Cost Assumptions", 24, ACCENT)
    table = slide.shapes.add_table(4, 3, Inches(0.85), Inches(1.45), Inches(9.4), Inches(2.2)).table
    table.cell(0, 0).text = "Environment"
    table.cell(0, 1).text = "Estimate"
    table.cell(0, 2).text = "Notes"
    add_text(slide, 0.9, 4.2, 10.4, 0.5, "Cost note", 16, MUTED)
    ensure_shapes(slide, 3)

    # 12 component cards.
    slide = prs.slides.add_slide(blank)
    set_bg(slide)
    add_text(slide, 0.7, 0.55, 5.0, 0.4, "Core Components", 24, ACCENT)
    for index in range(4):
        x = 0.8 + (index % 2) * 6.0
        y = 1.35 + (index // 2) * 2.35
        add_box(slide, x, y, 5.3, 1.7, "", fill=PANEL)
        add_text(slide, x + 0.25, y + 0.25, 4.8, 0.35, "Service", 18)
        add_text(slide, x + 0.25, y + 0.72, 4.8, 0.65, "Purpose", 14, MUTED)
    ensure_shapes(slide, 21)

    # 13 logical architecture labels.
    slide = prs.slides.add_slide(blank)
    set_bg(slide)
    add_text(slide, 0.7, 0.55, 5.0, 0.4, "Logical Architecture", 24, ACCENT)
    for index in range(8):
        x = 0.7 + (index % 4) * 3.15
        y = 1.4 + (index // 4) * 2.25
        add_box(slide, x, y, 2.45, 1.15, f"Node {index + 1}", fill=PANEL, font_size=15)
        add_text(slide, x, y + 1.28, 2.45, 0.2, "", 8)
    ensure_shapes(slide, 17)

    # 14 risks.
    slide = prs.slides.add_slide(blank)
    set_bg(slide)
    add_text(slide, 0.7, 0.55, 5.0, 0.4, "Risks", 24, ACCENT)
    for index in range(3):
        x = 0.8 + index * 4.15
        add_box(slide, x, 1.55, 3.65, 4.4, "", fill=PANEL)
        add_box(slide, x + 0.18, 1.78, 0.5, 0.5, "!", fill=ACCENT, line=ACCENT)
        add_text(slide, x + 0.25, 2.55, 3.0, 0.45, "Risk title", 18)
        add_text(slide, x + 0.25, 3.1, 3.0, 1.6, "Mitigation", 14, MUTED)
    ensure_shapes(slide, 16)

    # 15 next steps.
    slide = prs.slides.add_slide(blank)
    set_bg(slide)
    add_text(slide, 0.7, 0.55, 5.0, 0.4, "Next Steps", 24, ACCENT)
    for index in range(4):
        x = 0.8 + (index % 2) * 6.0
        y = 1.35 + (index // 2) * 2.35
        add_box(slide, x, y, 5.3, 1.7, "", fill=PANEL)
        add_text(slide, x + 0.25, y + 0.25, 4.8, 0.35, "Step title", 18)
        add_text(slide, x + 0.25, y + 0.72, 4.8, 0.65, "Step body", 14, MUTED)
    ensure_shapes(slide, 18)

    # 16 closing.
    slide = prs.slides.add_slide(blank)
    set_bg(slide)
    add_box(slide, 0.0, 0.0, 13.333, 0.16, "", fill=ACCENT, line=ACCENT)
    add_text(slide, 1.0, 2.2, 7.0, 0.8, "Thank You", 40)
    add_text(slide, 1.0, 3.25, 8.8, 0.7, "Closing note", 20, MUTED)
    add_text(slide, 1.0, 6.35, 6.0, 0.3, "AI Proposal Workbench", 14, MUTED)
    ensure_shapes(slide, 4)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_template()
