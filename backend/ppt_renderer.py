import json
import sys
import uuid
from pathlib import Path
from typing import Any, Dict, Iterable, List, Sequence, Tuple

CURRENT_DIR = Path(__file__).resolve().parent
if str(CURRENT_DIR) not in sys.path:
    sys.path.insert(0, str(CURRENT_DIR))

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import MSO_VERTICAL_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

from proposal_agent import proposal_output_dir


ROOT_DIR = Path(__file__).resolve().parent.parent
TEMPLATE_DIR = ROOT_DIR / "templates"
REGISTRY_PATH = TEMPLATE_DIR / "template_registry.json"

BG = RGBColor(246, 248, 251)
CARD = RGBColor(255, 255, 255)
TEXT = RGBColor(15, 23, 42)
MUTED = RGBColor(71, 85, 105)
SOFT = RGBColor(148, 163, 184)
BORDER = RGBColor(224, 229, 236)
ACCENT = RGBColor(20, 184, 166)
ACCENT_2 = RGBColor(37, 99, 235)
ACCENT_3 = RGBColor(245, 158, 11)
ACCENT_4 = RGBColor(168, 85, 247)


def load_registry() -> Dict[str, Any]:
    if not REGISTRY_PATH.exists():
        return {"templates": {}}
    return json.loads(REGISTRY_PATH.read_text(encoding="utf-8"))


def list_templates() -> Dict[str, Any]:
    return load_registry().get("templates", {})


def _flatten(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, str):
        return value
    if isinstance(value, list):
        if all(isinstance(item, dict) for item in value):
            rows = []
            for item in value:
                rows.append("; ".join(f"{key}: {val}" for key, val in item.items()))
            return "\n".join(rows)
        return "\n".join(f"- {item}" for item in value)
    if isinstance(value, dict):
        return "\n".join(f"{key}: {_flatten(val)}" for key, val in value.items())
    return str(value)


def _section_value(proposal: Dict[str, Any], section_key: str) -> str:
    return _flatten(proposal.get(section_key, ""))


def _set_background(slide) -> None:
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = BG


def _shape_rect(slide, left, top, width, height, fill=CARD, line=BORDER, radius=True):
    shape_type = MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE if radius else MSO_AUTO_SHAPE_TYPE.RECTANGLE
    shape = slide.shapes.add_shape(shape_type, Inches(left), Inches(top), Inches(width), Inches(height))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = line
    shape.line.width = Pt(1)
    return shape


def _shape_line(slide, left, top, width, height=0.02, color=BORDER):
    shape = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height))
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    return shape


def _textbox(slide, left, top, width, height):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.06)
    tf.margin_right = Inches(0.06)
    tf.margin_top = Inches(0.04)
    tf.margin_bottom = Inches(0.04)
    return box, tf


def _apply_text_style(paragraph, size=16, color=TEXT, bold=False, align=PP_ALIGN.LEFT, font_name="Aptos"):
    paragraph.alignment = align
    paragraph.font.size = Pt(size)
    paragraph.font.color.rgb = color
    paragraph.font.bold = bold
    paragraph.font.name = font_name


def _set_shape_text(shape, lines: Sequence[str], title_size=16, body_size=14, title_color=TEXT, body_color=MUTED):
    tf = shape.text_frame
    tf.clear()
    if not lines:
        return
    first = tf.paragraphs[0]
    first.text = lines[0]
    _apply_text_style(first, size=title_size, color=title_color, bold=True)
    if len(lines) > 1:
        for line in lines[1:]:
            p = tf.add_paragraph()
            p.text = line
            _apply_text_style(p, size=body_size, color=body_color)
            p.space_after = Pt(2)


def _set_textbox(tf, title: str, body: str, title_size=18, body_size=15, body_color=MUTED, title_color=TEXT):
    tf.clear()
    p = tf.paragraphs[0]
    p.text = title
    _apply_text_style(p, size=title_size, color=title_color, bold=True)
    p.space_after = Pt(6)
    for line in [line.strip() for line in body.split("\n") if line.strip()]:
        p = tf.add_paragraph()
        p.text = line
        _apply_text_style(p, size=body_size, color=body_color)
        p.space_after = Pt(3)


def _add_footer(slide, proposal: Dict[str, Any], template_id: str, page_label: str):
    _shape_line(slide, 0.55, 7.02, 12.2, 0.015, color=BORDER)
    _, tf = _textbox(slide, 0.7, 7.06, 6.6, 0.18)
    tf.text = f"{page_label}  |  Template: {template_id}"
    _apply_text_style(tf.paragraphs[0], size=9, color=SOFT)

    _, right_tf = _textbox(slide, 8.8, 7.06, 3.7, 0.18)
    right_tf.text = proposal.get("generated_at", "")
    _apply_text_style(right_tf.paragraphs[0], size=9, color=SOFT, align=PP_ALIGN.RIGHT)


def _add_section_title(slide, title: str, subtitle: str = ""):
    _shape_line(slide, 0.7, 0.52, 1.15, 0.07, color=ACCENT)
    _, tf = _textbox(slide, 0.7, 0.64, 10.8, 0.56)
    tf.text = title
    _apply_text_style(tf.paragraphs[0], size=30, color=TEXT, bold=True)
    if subtitle:
        _, sub_tf = _textbox(slide, 0.72, 1.08, 11.3, 0.46)
        sub_tf.text = subtitle
        _apply_text_style(sub_tf.paragraphs[0], size=13, color=MUTED)


def _add_bullet_box(slide, left, top, width, height, title, bullets, accent=ACCENT_2):
    card = _shape_rect(slide, left, top, width, height)
    _shape_line(slide, left, top, width, 0.08, color=accent)
    _, tf = _textbox(slide, left + 0.18, top + 0.14, width - 0.36, height - 0.24)
    tf.clear()
    p = tf.paragraphs[0]
    p.text = title
    _apply_text_style(p, size=17, color=TEXT, bold=True)
    for bullet in bullets:
        p = tf.add_paragraph()
        p.text = bullet
        _apply_text_style(p, size=14, color=MUTED)
        p.level = 0
        p.space_after = Pt(5)
    return card


def _add_mini_card(slide, left, top, width, height, title, body, accent=ACCENT_2):
    card = _shape_rect(slide, left, top, width, height)
    _shape_line(slide, left, top, width, 0.07, color=accent)
    _, tf = _textbox(slide, left + 0.15, top + 0.14, width - 0.3, height - 0.22)
    tf.clear()
    p = tf.paragraphs[0]
    p.text = title
    _apply_text_style(p, size=15, color=TEXT, bold=True)
    p.space_after = Pt(4)
    for line in body.split("\n"):
        if not line.strip():
            continue
        p = tf.add_paragraph()
        p.text = line.strip()
        _apply_text_style(p, size=12.5, color=MUTED)
        p.space_after = Pt(2)
    return card


def _add_cover_slide(prs: Presentation, proposal: Dict[str, Any], template_id: str):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _set_background(slide)
    _shape_line(slide, 0.0, 0.0, 13.333, 0.14, color=ACCENT_2)

    # Left side: title and purpose
    _, label_tf = _textbox(slide, 0.82, 0.75, 2.5, 0.28)
    label_tf.text = "PROPOSAL GENERATOR"
    _apply_text_style(label_tf.paragraphs[0], size=11, color=ACCENT_2, bold=True)

    _, tf = _textbox(slide, 0.82, 1.02, 6.3, 1.45)
    tf.text = "AI Solution Architect\nProposal"
    for paragraph in tf.paragraphs:
        _apply_text_style(paragraph, size=33, color=TEXT, bold=True)
        paragraph.space_after = Pt(0)

    _, sub_tf = _textbox(slide, 0.84, 2.55, 6.0, 0.9)
    sub_tf.text = (
        "Turn an uploaded requirement document into structured requirements, clarification questions, "
        "an implementable architecture, and a customer-ready PowerPoint deck."
    )
    _apply_text_style(sub_tf.paragraphs[0], size=16, color=MUTED)

    # Right side: a cleaner summary panel
    panel = _shape_rect(slide, 7.95, 0.86, 4.35, 4.36)
    _shape_line(slide, 7.95, 0.86, 4.35, 0.09, color=ACCENT)
    _, summary_tf = _textbox(slide, 8.18, 1.08, 3.85, 0.48)
    summary_tf.text = "What this deck delivers"
    _apply_text_style(summary_tf.paragraphs[0], size=17, color=TEXT, bold=True)
    summary_items = [
        "A concise executive summary.",
        "Expected output and engineering challenges.",
        "A realistic solution blueprint and stack.",
        "Environment-level cost framing.",
        "A clear delivery roadmap and risk profile.",
    ]
    y = 1.62
    for idx, item in enumerate(summary_items):
        pill = _shape_rect(slide, 8.18, y, 3.78, 0.48, fill=BG, line=BORDER, radius=True)
        _, ptf = _textbox(slide, 8.38, y + 0.05, 3.45, 0.25)
        ptf.text = item
        _apply_text_style(ptf.paragraphs[0], size=12.5, color=MUTED)
        y += 0.66

    # Bottom strip: simple visual flow
    flow_y = 4.98
    flow_cards = [
        ("1", "Extract", "Pull key requirements, constraints, and roles."),
        ("2", "Clarify", "Ask focused follow-up questions when details are missing."),
        ("3", "Export", "Render the final proposal into the selected template."),
    ]
    flow_x = [0.84, 3.18, 5.52]
    for (num, title, body), left in zip(flow_cards, flow_x):
        card = _shape_rect(slide, left, flow_y, 2.05, 1.18)
        badge = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.OVAL, Inches(left + 0.15), Inches(flow_y + 0.16), Inches(0.44), Inches(0.44))
        badge.fill.solid()
        badge.fill.fore_color.rgb = ACCENT_2 if num == "1" else ACCENT_3 if num == "2" else ACCENT_4
        badge.line.fill.background()
        _, ntf = _textbox(slide, left + 0.15, flow_y + 0.2, 0.44, 0.24)
        ntf.text = num
        _apply_text_style(ntf.paragraphs[0], size=13, color=RGBColor(255, 255, 255), bold=True, align=PP_ALIGN.CENTER)
        _, card_tf = _textbox(slide, left + 0.67, flow_y + 0.12, 1.2, 0.85)
        card_tf.text = title
        _apply_text_style(card_tf.paragraphs[0], size=15, color=TEXT, bold=True)
        p = card_tf.add_paragraph()
        p.text = body
        _apply_text_style(p, size=11.5, color=MUTED)
    arrow1 = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.CHEVRON, Inches(2.88), Inches(flow_y + 0.38), Inches(0.18), Inches(0.18))
    arrow1.fill.solid()
    arrow1.fill.fore_color.rgb = BORDER
    arrow1.line.fill.background()
    arrow2 = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.CHEVRON, Inches(5.22), Inches(flow_y + 0.38), Inches(0.18), Inches(0.18))
    arrow2.fill.solid()
    arrow2.fill.fore_color.rgb = BORDER
    arrow2.line.fill.background()

    _, tf2 = _textbox(slide, 8.18, 5.48, 3.8, 0.3)
    tf2.text = f"Template-ready export  |  {template_id}"
    _apply_text_style(tf2.paragraphs[0], size=10.5, color=SOFT)

    _add_footer(slide, proposal, template_id, "Cover")
    return slide


def _add_expected_output_slide(prs: Presentation, proposal: Dict[str, Any], template_id: str):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _set_background(slide)
    _add_section_title(
        slide,
        "Expected Output",
        "The output should read like a real proposal, not a raw AI transcript.",
    )

    _add_bullet_box(
        slide,
        0.72,
        1.72,
        7.2,
        4.55,
        "Required outcomes",
        proposal.get("expected_output", [
            "Extract the key requirements from the uploaded document.",
            "Ask clarification questions when details are missing.",
            "Propose an implementable AI architecture and deployment design.",
            "Estimate costs for Development, Test, and UAT or Staging environments.",
            "Export a customer-ready PowerPoint proposal using the selected enterprise template.",
        ]),
        accent=ACCENT_2,
    )

    _add_mini_card(
        slide,
        8.2,
        1.72,
        3.9,
        1.28,
        "Customer value",
        "Faster proposal drafting, clearer review cycles, and a cleaner handoff to delivery teams.",
        accent=ACCENT,
    )
    _add_mini_card(
        slide,
        8.2,
        3.14,
        3.9,
        1.28,
        "HITL loop",
        "Users can refine assumptions, architecture direction, and final slide wording before export.",
        accent=ACCENT_3,
    )
    _add_mini_card(
        slide,
        8.2,
        4.56,
        3.9,
        1.28,
        "Template discipline",
        "All generated content should fit the selected enterprise layout without breaking branding.",
        accent=ACCENT_4,
    )

    _add_footer(slide, proposal, template_id, "Expected Output")
    return slide


def _challenge_cards() -> List[Tuple[str, str, RGBColor]]:
    return [
        ("Understand the brief", "Interpret messy, unstructured requirement documents accurately.", ACCENT),
        ("Avoid hallucination", "Keep recommendations specific, implementable, and not overly generic.", ACCENT_2),
        ("Stay consistent", "Keep business, technical, infrastructure, and cost sections aligned.", ACCENT_3),
        ("Map into slides", "Fit AI-generated content into PPT layouts without breaking the design.", ACCENT_4),
        ("Support templates", "Handle multiple proposal templates and future branding changes.", ACCENT),
        ("Preserve HITL", "Let reviewers refine assumptions, solution direction, and slide content.", ACCENT_2),
    ]


def _add_challenges_slide(prs: Presentation, proposal: Dict[str, Any], template_id: str):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _set_background(slide)
    _add_section_title(
        slide,
        "Key Engineering Challenges",
        "These are the parts that make the proposal system trustworthy instead of gimmicky.",
    )

    cards = _challenge_cards()
    positions = [
        (0.72, 1.75), (4.26, 1.75), (7.8, 1.75),
        (0.72, 3.32), (4.26, 3.32), (7.8, 3.32),
    ]
    for (title, body, accent), (left, top) in zip(cards, positions):
        _add_mini_card(slide, left, top, 3.1, 1.18, title, body, accent=accent)

    _add_footer(slide, proposal, template_id, "Challenges")
    return slide


def _add_blueprint_slide(prs: Presentation, proposal: Dict[str, Any], template_id: str):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _set_background(slide)
    _add_section_title(
        slide,
        "Proposed Solution Blueprint",
        "A compact flow keeps the narrative clear from intake to export.",
    )

    steps = [
        ("1", "Ingest", "Upload PDF, DOCX, TXT, or MD and normalize the content."),
        ("2", "Extract", "Pull structured requirements and flag any missing details."),
        ("3", "Draft", "Generate architecture, proposal sections, risks, and assumptions."),
        ("4", "Render", "Map the draft into the selected PPT template and export PPTX."),
    ]
    lefts = [0.72, 3.95, 7.18, 10.41]
    for idx, ((num, title, body), left) in enumerate(zip(steps, lefts)):
        card = _shape_rect(slide, left, 1.98, 2.15, 2.32)
        _shape_line(slide, left, 1.98, 2.15, 0.08, color=[ACCENT, ACCENT_2, ACCENT_3, ACCENT_4][idx])
        circle = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.OVAL, Inches(left + 0.18), Inches(2.18), Inches(0.52), Inches(0.52))
        circle.fill.solid()
        circle.fill.fore_color.rgb = [ACCENT, ACCENT_2, ACCENT_3, ACCENT_4][idx]
        circle.line.fill.background()
        _, ntf = _textbox(slide, left + 0.18, 2.24, 0.52, 0.3)
        ntf.text = num
        _apply_text_style(ntf.paragraphs[0], size=16, color=RGBColor(255, 255, 255), bold=True, align=PP_ALIGN.CENTER)
        _, tf = _textbox(slide, left + 0.16, 2.86, 1.82, 1.12)
        tf.text = title
        _apply_text_style(tf.paragraphs[0], size=16, color=TEXT, bold=True)
        p = tf.add_paragraph()
        p.text = body
        _apply_text_style(p, size=12.5, color=MUTED)
        p.space_before = Pt(5)

    _shape_line(slide, 1.79, 3.02, 9.63, 0.03, color=BORDER)
    _add_footer(slide, proposal, template_id, "Blueprint")
    return slide


def _add_stack_slide(prs: Presentation, proposal: Dict[str, Any], template_id: str):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _set_background(slide)
    _add_section_title(
        slide,
        "Recommended Stack",
        "A cloud-neutral stack keeps the demo free-friendly and still production-shaped.",
    )

    stack_groups = [
        (
            "AI & reasoning",
            [
                "Gemini API or local LLM fallback",
                "Proposal agent for extraction and drafting",
                "Human-in-the-loop clarification review",
            ],
            ACCENT,
        ),
        (
            "Retrieval & storage",
            [
                "FAISS, Chroma, or pgvector",
                "Local disk or S3-compatible storage",
                "PostgreSQL for projects and revisions",
            ],
            ACCENT_2,
        ),
        (
            "Delivery & governance",
            [
                "Template-driven PPTX renderer",
                "Optional Tavily / SerpAPI research tools",
                "Telemetry for quality and traceability",
            ],
            ACCENT_3,
        ),
    ]

    lefts = [0.72, 4.46, 8.2]
    for (title, bullets, accent), left in zip(stack_groups, lefts):
        _add_bullet_box(slide, left, 1.8, 3.0, 4.35, title, bullets, accent=accent)

    _add_footer(slide, proposal, template_id, "Stack")
    return slide


def _add_cost_slide(prs: Presentation, proposal: Dict[str, Any], template_id: str):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _set_background(slide)
    _add_section_title(
        slide,
        "Cost Estimation",
        "These are planning numbers for a lean free-friendly stack, not a rigid quote.",
    )

    costs = proposal.get("cost_estimation", [])
    cols = [0.72, 4.47, 8.22]
    accents = [ACCENT, ACCENT_2, ACCENT_3]
    for idx, (entry, left) in enumerate(zip(costs, cols)):
        card = _shape_rect(slide, left, 1.84, 3.05, 2.85)
        _shape_line(slide, left, 1.84, 3.05, 0.08, color=accents[idx])
        _, tf = _textbox(slide, left + 0.18, 2.08, 2.68, 2.4)
        tf.clear()
        p = tf.paragraphs[0]
        p.text = entry.get("environment", "Environment")
        _apply_text_style(p, size=18, color=TEXT, bold=True)
        p.space_after = Pt(8)
        p = tf.add_paragraph()
        p.text = entry.get("monthly_estimate_usd", "")
        _apply_text_style(p, size=24, color=accents[idx], bold=True)
        p.space_after = Pt(8)
        p = tf.add_paragraph()
        p.text = entry.get("notes", "")
        _apply_text_style(p, size=12.5, color=MUTED)

    note = (
        "If the deployment stays local-first, the Development and Test environments can be kept relatively light. "
        "When external research or hosted storage is added, the cost profile increases with usage."
    )
    _add_mini_card(slide, 0.72, 5.1, 10.55, 1.08, "Cost note", note, accent=ACCENT_4)

    _add_footer(slide, proposal, template_id, "Cost")
    return slide


def _add_delivery_slide(prs: Presentation, proposal: Dict[str, Any], template_id: str):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _set_background(slide)
    _add_section_title(
        slide,
        "Delivery Plan and Risks",
        "A short delivery path keeps the scope understandable for demo and production planning.",
    )

    phases = proposal.get("implementation_approach", [])
    risks = proposal.get("risks", [])

    _add_bullet_box(
        slide,
        0.72,
        1.78,
        5.9,
        4.6,
        "Implementation approach",
        phases,
        accent=ACCENT,
    )

    _add_bullet_box(
        slide,
        6.95,
        1.78,
        5.65,
        4.6,
        "Risks and mitigations",
        risks,
        accent=ACCENT_3,
    )

    _add_footer(slide, proposal, template_id, "Delivery")
    return slide


def _compact_slide_text(value: Any, max_length: int = 155) -> str:
    text = _clean_ppt_text(" ".join(str(value).split()))
    if len(text) <= max_length:
        return text
    shortened = text[: max_length - 3].rsplit(" ", 1)[0]
    return f"{shortened}..."


def _clean_ppt_text(value: Any) -> str:
    text = "" if value is None else str(value)
    replacements = {
        "â€”": "-",
        "â€“": "-",
        "â†’": "->",
        "â‰¤": "<=",
        "â€‘": "-",
        "â€œ": '"',
        "â€": '"',
        "â€™": "'",
        "\u2014": "-",
        "\u2013": "-",
        "\u2192": "->",
        "\u2264": "<=",
        "\u2011": "-",
    }
    for bad, good in replacements.items():
        text = text.replace(bad, good)
    return text


def _pptagent_slides(proposal: Dict[str, Any]) -> List[Dict[str, Any]]:
    pptagent = proposal.get("pptagent_input") or {}
    slides = pptagent.get("slides") if isinstance(pptagent, dict) else []
    return [slide for slide in slides or [] if isinstance(slide, dict)]


def _slide_by_type(agent_slides: List[Dict[str, Any]], *slide_types: str) -> Dict[str, Any]:
    wanted = set(slide_types)
    return next((slide for slide in agent_slides if slide.get("slide_type") in wanted or slide.get("layout_id") in wanted), {})


def _slide_title(slide: Dict[str, Any], fallback: str = "") -> str:
    return _compact_slide_text(slide.get("title") or slide.get("primary_message") or fallback, 90)


def _slide_subtitle(slide: Dict[str, Any], fallback: str = "") -> str:
    return _compact_slide_text(slide.get("subtitle") or slide.get("primary_message") or fallback, 120)


def _slide_content(slide: Dict[str, Any]) -> Dict[str, Any]:
    content = slide.get("content") or {}
    return content if isinstance(content, dict) else {}


def _slide_bullets(slide: Dict[str, Any], limit: int = 5, max_length: int = 110) -> List[str]:
    bullets = _slide_content(slide).get("bullets") or []
    return [_compact_slide_text(item, max_length) for item in bullets if str(item).strip()][:limit]


def _slide_kv(slide: Dict[str, Any], limit: int = 4) -> List[Tuple[str, str]]:
    pairs = []
    for item in (_slide_content(slide).get("key_value_pairs") or [])[:limit]:
        if isinstance(item, dict):
            pairs.append((_compact_slide_text(item.get("key"), 35), _compact_slide_text(item.get("value"), 80)))
    return pairs


def _slide_steps(slide: Dict[str, Any], limit: int = 5) -> List[Tuple[str, str]]:
    steps = []
    for item in (_slide_content(slide).get("process_steps") or [])[:limit]:
        if isinstance(item, dict):
            label = item.get("label") or f"Step {item.get('step_number') or len(steps) + 1}"
            steps.append((_compact_slide_text(label, 32), _compact_slide_text(item.get("description"), 75)))
    return steps


def _slide_component_cards(slide: Dict[str, Any], limit: int = 4) -> List[Dict[str, str]]:
    cards = []
    for item in (_slide_content(slide).get("component_cards") or [])[:limit]:
        if isinstance(item, dict):
            cards.append(
                {
                    "service": _compact_slide_text(item.get("name") or item.get("service") or item.get("title"), 35),
                    "purpose": _compact_slide_text(item.get("description") or item.get("purpose"), 75),
                }
            )
    return cards


def _proposal_slide_plan(proposal: Dict[str, Any]) -> List[Dict[str, Any]]:
    plan = proposal.get("slide_plan", [])
    valid = []
    for item in plan:
        if not isinstance(item, dict) or not item.get("title"):
            continue
        bullets = [_compact_slide_text(bullet) for bullet in item.get("bullets", []) if str(bullet).strip()]
        valid.append({"title": str(item["title"]), "bullets": bullets[:5]})
    if valid:
        return valid

    return [
        {
            "title": title,
            "bullets": [_compact_slide_text(item) for item in proposal.get(section, [])][:5]
            if isinstance(proposal.get(section), list)
            else [_compact_slide_text(proposal.get(section, ""))],
        }
        for section, title in PROPOSAL_SECTIONS
    ]


def _add_proposal_section_slide(
    prs: Presentation,
    proposal: Dict[str, Any],
    template_id: str,
    page_number: int,
    title: str,
    bullets: List[str],
):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _set_background(slide)
    _add_section_title(slide, title, "Approved proposal content generated from the reviewed requirements and architecture.")

    accent = [ACCENT, ACCENT_2, ACCENT_3, ACCENT_4][page_number % 4]
    _add_bullet_box(slide, 0.72, 1.72, 7.55, 4.95, "Key points", bullets or ["No approved content was provided for this section."], accent=accent)

    decisions = proposal.get("architecture_decisions", {})
    decision_summary = decisions.get("executive_recommendation", "") if isinstance(decisions, dict) else ""
    context = proposal.get("problem_statement") or proposal.get("executive_summary") or ""
    _add_mini_card(slide, 8.55, 1.72, 3.78, 2.12, "Proposal context", _compact_slide_text(context, 280), accent=ACCENT_2)
    _add_mini_card(
        slide,
        8.55,
        4.1,
        3.78,
        2.0,
        "Architecture decision",
        _compact_slide_text(decision_summary, 260) or "Architecture approval is required before final export.",
        accent=ACCENT_3,
    )
    _add_footer(slide, proposal, template_id, f"{page_number:02d}")
    return slide


def _output_path(template_id: str, proposal_id: str) -> Path:
    output_dir = proposal_output_dir()
    output = output_dir / f"proposal_{template_id}_{proposal_id}.pptx"
    if not output.exists():
        return output
    return output_dir / f"proposal_{template_id}_{proposal_id}_{uuid.uuid4().hex[:8]}.pptx"


def render_from_template(proposal: Dict[str, Any], template_id: str) -> Path:
    templates = list_templates()
    config = templates.get(template_id)
    if not config:
        raise ValueError(f"Unknown template_id '{template_id}'.")

    template_path = ROOT_DIR / config.get("template_path", "")
    if not template_path.exists():
        return render_starter_deck(proposal, template_id)

    prs = Presentation(str(template_path))
    sections = config.get("sections", {})
    replacements = {key: _section_value(proposal, key) for key in sections.keys()}

    for slide in prs.slides:
        for shape in slide.shapes:
            if hasattr(shape, "text_frame") and shape.text_frame is not None:
                original = shape.text
                updated = original
                for key, value in replacements.items():
                    updated = updated.replace("{{" + key + "}}", value)
                if updated != original:
                    shape.text = updated

    output = _output_path(template_id, proposal["proposal_id"])
    prs.save(output)
    return output


def _master_template_agenda(proposal: Dict[str, Any]) -> List[str]:
    agent_slides = _pptagent_slides(proposal)
    agenda = [_slide_title(slide, f"Section {index}") for index, slide in enumerate(agent_slides[1:8], start=1)]
    if agenda:
        return agenda[:7]
    return [
        "Business context and objectives",
        "Solution overview",
        "Core solution components",
        "Logical architecture",
        "Cost assumptions",
        "Risks and mitigations",
        "Recommended next steps",
    ]


def _master_template_kpis(proposal: Dict[str, Any]) -> List[Tuple[str, str, str]]:
    agent_slides = _pptagent_slides(proposal)
    goals = _slide_by_type(agent_slides, "goals_kpi")
    pairs = _slide_kv(goals, 3)
    if pairs:
        return [(key, "confirmed", value) for key, value in pairs[:3]]
    expected = proposal.get("expected_output", []) or []
    risks = proposal.get("risks", []) or []
    roadmap = proposal.get("implementation_approach", []) or []
    return [
        ("POC", "ready to validate", f"{min(len(expected), 5) or 3} validated outputs"),
        ("Low-risk", "needs review", f"{min(len(risks), 3) or 3} key mitigations"),
        ("Delivery", "structured", f"{min(len(roadmap), 4) or 4} rollout steps"),
    ]


def _master_template_cost_rows(proposal: Dict[str, Any]) -> List[List[str]]:
    rows: List[List[str]] = [["Environment", "Estimate (USD/mo)", "Notes"]]
    agent_slides = _pptagent_slides(proposal)
    cost_slide = _slide_by_type(agent_slides, "cost_table", "cost", "cost_timeline")
    table = _slide_content(cost_slide).get("table") or {}
    table_rows = table.get("rows") if isinstance(table, dict) else []
    if isinstance(table_rows, list) and table_rows:
        for item in table_rows[:3]:
            if isinstance(item, dict):
                rows.append(
                    [
                        _compact_slide_text(item.get("environment") or item.get("service") or "POC", 28),
                        _compact_slide_text(item.get("estimate") or item.get("monthly_estimate_usd") or item.get("cost") or "TBD", 24),
                        _compact_slide_text(item.get("notes") or item.get("assumption") or item.get("description") or "", 72),
                    ]
                )
            elif isinstance(item, list):
                rows.append([_compact_slide_text(value, 72) for value in item[:3]])
        if len(rows) > 1:
            return rows[:4]
    for item in proposal.get("cost_estimation", []) or []:
        if not isinstance(item, dict):
            continue
        rows.append(
            [
                str(item.get("environment") or "TBD"),
                str(item.get("monthly_estimate_usd") or "TBD"),
                _compact_slide_text(item.get("notes") or "", 72) or "Assumption pending",
            ]
        )
    if len(rows) == 1:
        rows.extend(
            [
                ["Development", "TBD", "Usage assumptions pending"],
                ["Test", "TBD", "Usage assumptions pending"],
                ["UAT / Staging", "TBD", "Usage assumptions pending"],
            ]
        )
    return rows[:4]


def _set_shape_text_direct(shape, text: str) -> None:
    if hasattr(shape, "text_frame") and shape.text_frame is not None:
        shape.text = text


def _set_slide_shape_text(slide, index: int, text: str) -> None:
    if index < len(slide.shapes):
        _set_shape_text_direct(slide.shapes[index], text)


def _style_shape_text(
    shape,
    *,
    size: int = 18,
    color=TEXT,
    bold: bool = False,
    uppercase: bool = False,
    align=PP_ALIGN.LEFT,
) -> None:
    if not hasattr(shape, "text_frame") or shape.text_frame is None:
        return
    for paragraph in shape.text_frame.paragraphs:
        if uppercase and paragraph.text:
            paragraph.text = paragraph.text.upper()
        paragraph.alignment = align
        for run in paragraph.runs:
            run.font.size = Pt(size)
            run.font.color.rgb = color
            run.font.bold = bold
            run.font.name = "Aptos"


def _replace_circle_with_header_strip(slide, circle_index: int, title_index: int, body_index: int, accent) -> None:
    if body_index >= len(slide.shapes):
        return
    if circle_index < len(slide.shapes):
        circle = slide.shapes[circle_index]
        circle.fill.background()
        circle.line.fill.background()
        if hasattr(circle, "text_frame") and circle.text_frame is not None:
            circle.text = ""
    if title_index < len(slide.shapes):
        _style_shape_text(slide.shapes[title_index], size=20, color=accent, bold=True, uppercase=True)
    if body_index < len(slide.shapes):
        _style_shape_text(slide.shapes[body_index], size=18, color=TEXT, bold=False)


def _render_master_template(proposal: Dict[str, Any], template_id: str = "master_template") -> Path:
    templates = list_templates()
    config = templates.get(template_id)
    if not config:
        raise ValueError(f"Unknown template_id '{template_id}'.")

    template_path = ROOT_DIR / config.get("template_path", "")
    if not template_path.exists():
        raise FileNotFoundError(f"Master template not found: {template_path}")

    prs = Presentation(str(template_path))
    slides = prs.slides
    agent_slides = _pptagent_slides(proposal)
    title_slide = _slide_by_type(agent_slides, "title")
    context_slide = _slide_by_type(agent_slides, "customer_context")
    problem_slide = _slide_by_type(agent_slides, "problem_impact")
    goals_slide = _slide_by_type(agent_slides, "goals_kpi")
    capability_slide = _slide_by_type(agent_slides, "capability_summary", "four_column_value_chain", "solution_overview")
    workflow_slide = _slide_by_type(agent_slides, "process_flow")
    architecture_slide = _slide_by_type(agent_slides, "logical_architecture", "high_level_architecture")
    components_slide = _slide_by_type(agent_slides, "component_cards")
    security_slide = _slide_by_type(agent_slides, "infrastructure_security", "security_operations")
    cost_slide = _slide_by_type(agent_slides, "cost_table", "cost", "cost_timeline")
    roadmap_slide = _slide_by_type(agent_slides, "roadmap", "implementation_roadmap", "next_steps")

    deck_title = _slide_title(title_slide, proposal.get("problem_statement") or proposal.get("executive_summary") or "Enterprise proposal")
    subtitle = _slide_subtitle(title_slide, proposal.get("executive_summary") or "AI proposal generation workbench")
    generated_date = proposal.get("generated_at", "")[:10] or "2026-07-26"

    # Slide 1 - cover
    _set_slide_shape_text(slides[0], 3, deck_title)
    _set_slide_shape_text(slides[0], 4, subtitle)
    _set_slide_shape_text(slides[0], 5, f"Prepared by AI Proposal Workbench      {generated_date}")
    _style_shape_text(slides[0].shapes[2], size=18, color=ACCENT, bold=True, uppercase=True)
    slides[0].shapes[3].height = Inches(0.8)
    slides[0].shapes[4].top = Inches(4.22)
    slides[0].shapes[4].height = Inches(0.55)
    _style_shape_text(slides[0].shapes[3], size=30, color=RGBColor(255, 255, 255), bold=True)
    _style_shape_text(slides[0].shapes[4], size=18, color=RGBColor(255, 255, 255), bold=False)
    _style_shape_text(slides[0].shapes[5], size=18, color=RGBColor(255, 255, 255))

    # Slide 2 - section opener
    _set_slide_shape_text(slides[1], 3, _slide_title(capability_slide, "Solution Overview"))
    _set_slide_shape_text(slides[1], 4, _compact_slide_text(subtitle, 90))
    _style_shape_text(slides[1].shapes[1], size=20, color=RGBColor(255, 255, 255), bold=True)
    _style_shape_text(slides[1].shapes[2], size=18, color=ACCENT, bold=True, uppercase=True)
    _style_shape_text(slides[1].shapes[3], size=34, color=RGBColor(255, 255, 255), bold=True)
    _style_shape_text(slides[1].shapes[4], size=20, color=RGBColor(255, 255, 255))

    # Slide 3 - agenda
    for shape_index, agenda_text in zip((3, 6, 9, 12, 15, 18, 21), _master_template_agenda(proposal)):
        _set_slide_shape_text(slides[2], shape_index, agenda_text)

    # Slide 4 - executive summary cards
    executive_cards = [
        (_slide_title(context_slide, "Customer Context"), " | ".join(_slide_bullets(context_slide, 2, 90)) or _compact_slide_text(proposal.get("executive_summary"), 90)),
        (_slide_title(problem_slide, "Problems and Impact"), " | ".join(_slide_bullets(problem_slide, 2, 90)) or _compact_slide_text(proposal.get("problem_statement"), 90)),
        (_slide_title(goals_slide, "Goals and Success Criteria"), " | ".join(_slide_bullets(goals_slide, 2, 90)) or _compact_slide_text((proposal.get("implementation_approach") or ["TBD"])[0], 90)),
    ]
    for title_index, body_index, card in zip((4, 9, 14), (5, 10, 15), executive_cards):
        _set_slide_shape_text(slides[3], title_index, card[0])
        _set_slide_shape_text(slides[3], body_index, card[1])
    _replace_circle_with_header_strip(slides[3], 3, 4, 5, ACCENT_2)
    _replace_circle_with_header_strip(slides[3], 8, 9, 10, ACCENT)
    _replace_circle_with_header_strip(slides[3], 13, 14, 15, ACCENT_4)

    # Slide 5 - why this solution
    _set_slide_shape_text(slides[4], 0, _slide_title(capability_slide, "Why this solution"))
    _set_slide_shape_text(slides[4], 1, _slide_subtitle(capability_slide, _slide_title(problem_slide, "")))
    principles = _slide_bullets(capability_slide, 4, 120) or (
        (proposal.get("business_solution") or [])[:2]
        + (proposal.get("technical_solution") or [])[:1]
        + (proposal.get("infrastructure_design") or [])[:1]
    )
    _set_slide_shape_text(slides[4], 2, "\n".join(f"- {_compact_slide_text(item, 120)}" for item in principles[:4]))
    _style_shape_text(slides[4].shapes[0], size=22, color=ACCENT_2, bold=True, uppercase=True)
    _style_shape_text(slides[4].shapes[1], size=18, color=TEXT, bold=True)
    _style_shape_text(slides[4].shapes[2], size=18, color=TEXT)

    # Slide 6 - business / technical split
    _set_slide_shape_text(slides[5], 2, _slide_title(problem_slide, "Business challenge"))
    _set_slide_shape_text(
        slides[5], 3, "\n".join(f"- {_compact_slide_text(item, 80)}" for item in (_slide_bullets(problem_slide, 3, 80) or proposal.get("business_solution") or ["TBD"])[:3])
    )
    _set_slide_shape_text(slides[5], 4, _slide_title(capability_slide, "Technical solution"))
    _set_slide_shape_text(
        slides[5], 5, "\n".join(f"- {_compact_slide_text(item, 80)}" for item in (_slide_bullets(capability_slide, 3, 80) or proposal.get("technical_solution") or ["TBD"])[:3])
    )

    # Slide 7 - engagement snapshot
    snapshot_pairs = [
        *(_slide_kv(context_slide, 4) or []),
        ("Primary stakeholder", "Proposal managers, solution architects, administrators"),
        ("Success criteria", _compact_slide_text(_slide_bullets(goals_slide, 1, 70)[0] if _slide_bullets(goals_slide, 1, 70) else "Customer-ready proposal export", 55)),
    ]
    for title_index, body_index, card in zip((3, 7, 11, 15), (4, 8, 12, 16), snapshot_pairs):
        _set_slide_shape_text(slides[6], title_index, card[0])
        _set_slide_shape_text(slides[6], body_index, card[1])

    # Slide 8 - target outcomes
    for base_index, (metric, status, label) in zip((2, 6, 10), _master_template_kpis(proposal)):
        _set_slide_shape_text(slides[7], base_index, metric)
        _set_slide_shape_text(slides[7], base_index + 1, status)
        _set_slide_shape_text(slides[7], base_index + 2, label)

    # Slide 9 - workflow
    workflow_steps = [
        ("Ingest", "Receive PDF, DOCX, TXT, and enterprise inputs"),
        ("Parse", "Use native parsing and OCR only when needed"),
        ("Analyze", "Extract deduplicated requirements with traceability"),
        ("Clarify", "Collect HITL answers before architecture"),
        ("Generate", "Produce proposal content and export to PowerPoint"),
    ]
    workflow_steps = _slide_steps(workflow_slide, 5) or workflow_steps
    for title_index, body_index, card in zip((4, 8, 12, 16, 20), (5, 9, 13, 17, 21), workflow_steps):
        _set_slide_shape_text(slides[8], title_index, card[0])
        _set_slide_shape_text(slides[8], body_index, card[1])

    # Slide 10 - timeline
    timeline = _slide_bullets(roadmap_slide, 4, 70) or (proposal.get("implementation_approach") or [])[:4]
    while len(timeline) < 4:
        timeline.append("Validate the next delivery checkpoint")
    for title_index, body_index, text in zip((4, 7, 10, 13), (4, 7, 10, 13), timeline):
        label = slides[9].shapes[title_index - 1].text if title_index - 1 < len(slides[9].shapes) else "Wk"
        _set_slide_shape_text(slides[9], body_index, _compact_slide_text(text, 45))

    # Slide 11 - cost assumptions
    table_shape = slides[10].shapes[1] if len(slides[10].shapes) > 1 and getattr(slides[10].shapes[1], "has_table", False) else None
    if table_shape is not None:
        rows = _master_template_cost_rows(proposal)
        table = table_shape.table
        for r, row in enumerate(rows):
            if r >= len(table.rows):
                break
            for c, value in enumerate(row):
                if c >= len(table.columns):
                    break
                table.cell(r, c).text = value
    _set_slide_shape_text(slides[10], 2, "Model usage charges are reported separately from base environment estimates.")

    # Slide 12 - core components
    components = _slide_component_cards(components_slide, 4) or (proposal.get("architecture_decisions", {}) or {}).get("components", []) or []
    fallback_components = [
        {"service": "Document processing", "purpose": "Parse text, layout, and scanned content"},
        {"service": "Requirement agent", "purpose": "Extract structured requirements with traceability"},
        {"service": "Architecture agent", "purpose": "Select patterns and Azure services"},
        {"service": "PPTX renderer", "purpose": "Map approved content into the final deck"},
    ]
    component_cards = (components[:4] if components else fallback_components)
    for title_index, body_index, item in zip((4, 9, 14, 19), (5, 10, 15, 20), component_cards):
        _set_slide_shape_text(slides[11], title_index, _compact_slide_text(item.get("service") or item.get("name") or "Component", 35))
        _set_slide_shape_text(slides[11], body_index, _compact_slide_text(item.get("purpose") or item.get("description") or "Purpose pending", 70))

    # Slide 13 - logical architecture
    architecture_services = []
    diagram = _slide_content(architecture_slide).get("diagram") or {}
    if isinstance(diagram, dict):
        architecture_services = [
            _compact_slide_text(node.get("label") or node.get("name") or node.get("id"), 35)
            for node in diagram.get("nodes", [])
            if isinstance(node, dict)
        ]
    architecture_services = architecture_services or (proposal.get("architecture_decisions", {}) or {}).get("azure_services", []) or []
    architecture_labels = [
        "Input Layer",
        "Requirement Documents",
        "AI Services",
        architecture_services[0] if len(architecture_services) > 0 else "Document Processing Service",
        architecture_services[1] if len(architecture_services) > 1 else "Proposal Agent",
        "Output Layer",
        architecture_services[2] if len(architecture_services) > 2 else "Template Resolver",
        architecture_services[3] if len(architecture_services) > 3 else "PPTX Renderer",
    ]
    for shape_index, text in zip((2, 4, 6, 8, 10, 12, 14, 16), architecture_labels):
        _set_slide_shape_text(slides[12], shape_index, str(text).replace(" ", "\n", 1) if shape_index in (4, 8, 10, 14, 16) and " " in str(text) else str(text))

    # Slide 14 - risks
    risks = _slide_bullets(security_slide, 3, 78) or proposal.get("risks", []) or ["Risk review pending"]
    while len(risks) < 3:
        risks.append("Risk review pending")
    titles = ["Scope pressure", "Template drift", "Content quality"]
    mitigations = [*risks[:3]]
    for title_index, body_index, title, body in zip((4, 9, 14), (5, 10, 15), titles, mitigations):
        _set_slide_shape_text(slides[13], title_index, title)
        _set_slide_shape_text(slides[13], body_index, _compact_slide_text(body, 78))

    # Slide 15 - next steps
    next_steps = _slide_bullets(roadmap_slide, 4, 70) or (proposal.get("implementation_approach") or [])[:4]
    while len(next_steps) < 4:
        next_steps.append("Confirm the next delivery milestone")
    step_titles = ["Finalize requirements", "Approve architecture", "Render proposal", "Validate output"]
    for title_index, body_index, title, body in zip((4, 8, 12, 16), (5, 9, 13, 17), step_titles, next_steps):
        _set_slide_shape_text(slides[14], title_index, title)
        _set_slide_shape_text(slides[14], body_index, _compact_slide_text(body, 52))

    # Slide 16 - closing
    _set_slide_shape_text(slides[15], 1, "Thank You")
    _set_slide_shape_text(slides[15], 2, _compact_slide_text(_slide_title(roadmap_slide, "Please review the proposed architecture, delivery plan, and template output."), 100))
    _set_slide_shape_text(slides[15], 3, "AI Proposal Workbench | local demo")

    output = _output_path(template_id, proposal["proposal_id"])
    prs.save(output)
    return output


def render_starter_deck(proposal: Dict[str, Any], template_id: str = "starter") -> Path:
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    _add_cover_slide(prs, proposal, template_id)
    for page_number, section in enumerate(_proposal_slide_plan(proposal), start=1):
        _add_proposal_section_slide(
            prs,
            proposal,
            template_id,
            page_number,
            section["title"],
            section["bullets"],
        )

    output = _output_path(template_id, proposal["proposal_id"])
    prs.save(output)
    return output


def render_pptx(proposal: Dict[str, Any], template_id: str = "starter") -> Path:
    if template_id == "starter":
        return render_starter_deck(proposal, template_id)
    if template_id == "master_template":
        return _render_master_template(proposal, template_id)
    return render_from_template(proposal, template_id)
