import json
import os
import re
import urllib.error
import urllib.request
import uuid
from datetime import datetime
from pathlib import Path
from typing import Any, Dict, List, Optional

from dotenv import load_dotenv

from foundry_responses import create_json_response, is_configured as azure_openai_configured
from foundry_published_agents import is_configured as published_agent_configured, invoke as invoke_published_agent
from slide_planner import build_slide_plan

load_dotenv(Path(__file__).resolve().parent.parent / ".env")


REQUIREMENT_FIELDS = [
    "business_goals",
    "functional_requirements",
    "non_functional_requirements",
    "constraints",
    "assumptions",
    "integrations",
    "security_needs",
    "expected_users",
]

PROPOSAL_SECTIONS = [
    ("executive_summary", "Executive Summary"),
    ("problem_statement", "Customer Challenge"),
    ("business_solution", "Business Solution"),
    ("technical_solution", "Technical Solution"),
    ("infrastructure_design", "Architecture and Infrastructure"),
    ("cost_estimation", "Cost Assumptions"),
    ("implementation_approach", "Implementation Roadmap"),
    ("risks", "Risks and Mitigations"),
]


FIELD_QUALITY_GUIDE = {
    "business_goals": "Clear target business outcomes, measurable success criteria, and the main pain points being solved.",
    "functional_requirements": "User-visible capabilities, workflows, data operations, and acceptance behavior.",
    "non_functional_requirements": "Scale, latency, availability, performance, compliance, observability, and maintainability.",
    "constraints": "Budget, timeline, hosting, technology, data residency, security, procurement, or policy limits.",
    "assumptions": "Explicit decisions used when the source document is silent or ambiguous.",
    "integrations": "Identity providers, data sources, APIs, enterprise systems, and downstream channels.",
    "security_needs": "Authentication, authorization, RBAC, encryption, audit, private networking, and data leakage controls.",
    "expected_users": "User personas, user count, concurrency, roles, and access boundaries.",
}


AI_PATTERN_KB = [
    {
        "name": "RAG with metadata filtering",
        "purpose": "Answer questions over enterprise documents while enforcing document-level and user-level filters.",
        "use_when": "Knowledge is mostly unstructured documents and answers must cite source material.",
        "data": "PDF, Word, HTML, wiki pages, tickets, policies, manuals.",
    },
    {
        "name": "Knowledge graph",
        "purpose": "Represent relationships between entities, policies, products, systems, and decisions.",
        "use_when": "The solution depends on multi-hop reasoning or explainable entity relationships.",
        "data": "Structured master data, taxonomies, relationship-heavy documents, CMDB, CRM, ERP metadata.",
    },
    {
        "name": "Multi-agent orchestration",
        "purpose": "Split work across specialist agents for research, extraction, validation, architecture, and writing.",
        "use_when": "The workflow has multiple reviewable stages and benefits from tool-using specialists.",
        "data": "Mixed documents, web sources, templates, cost tables, architecture inventories.",
    },
    {
        "name": "Computer-use automation",
        "purpose": "Operate legacy user interfaces when no stable API exists.",
        "use_when": "The process requires repetitive work in desktop or web applications that cannot be integrated directly.",
        "data": "Screens, forms, exported files, application UI state.",
    },
    {
        "name": "Workflow automation",
        "purpose": "Coordinate deterministic approvals, notifications, data movement, and business rules.",
        "use_when": "The solution needs reliable handoffs between people, systems, and scheduled jobs.",
        "data": "Events, forms, APIs, queues, business process states.",
    },
]

ARCHITECTURE_PATTERN_NAMES = [
    "RAG with metadata filtering",
    "Knowledge graph",
    "Multi-agent orchestration",
    "Computer-use automation",
    "Workflow automation",
    "Asynchronous document processing",
]

OFFICIAL_AZURE_SOURCE_HOSTS = (
    "learn.microsoft.com",
    "azure.microsoft.com",
    "github.com/azure",
    "github.com/microsoft",
)

RESEARCH_MODES = {"off", "azure_official_only", "azure_github_papers"}

ARCHITECT_RESEARCH_POLICIES = {
    "off": {
        "mode": "off",
        "description": "No public web search. Use confirmed requirements and private references only.",
        "allowed_source_types": [],
        "allowed_hosts": [],
        "safe_web_queries": [],
    },
    "azure_official_only": {
        "mode": "azure_official_only",
        "description": "Use Microsoft first-party Azure documentation and Microsoft/Azure GitHub repositories only.",
        "allowed_source_types": ["azure_docs", "microsoft_github"],
        "allowed_hosts": ["learn.microsoft.com", "azure.microsoft.com", "github.com/Azure", "github.com/microsoft"],
        "safe_web_queries": [
            "site:learn.microsoft.com Azure AI Search metadata filters security",
            "site:learn.microsoft.com Azure AI Document Intelligence layout model scanned PDF tables",
            "site:learn.microsoft.com Azure Container Apps jobs scale to zero",
            "site:learn.microsoft.com Azure Retail Prices API",
        ],
    },
    "azure_github_papers": {
        "mode": "azure_github_papers",
        "description": "Use Azure official docs, Microsoft/Azure GitHub references, and public research papers for architecture evidence.",
        "allowed_source_types": ["azure_docs", "microsoft_github", "research_paper"],
        "allowed_hosts": [
            "learn.microsoft.com",
            "azure.microsoft.com",
            "github.com/Azure",
            "github.com/microsoft",
            "arxiv.org",
            "aclanthology.org",
            "openreview.net",
            "paperswithcode.com",
        ],
        "safe_web_queries": [
            "site:learn.microsoft.com Azure AI Search metadata filters security",
            "site:learn.microsoft.com Azure AI Document Intelligence layout model scanned PDF tables",
            "site:learn.microsoft.com Azure Container Apps jobs scale to zero",
            "site:learn.microsoft.com Azure Retail Prices API",
            "site:github.com/Azure azure ai search samples vector search metadata filters",
            "site:github.com/microsoft document intelligence samples layout extraction",
            "site:arxiv.org retrieval augmented generation enterprise access control metadata filtering",
            "site:aclanthology.org retrieval augmented generation citation grounded answers",
            "site:openreview.net agentic workflow tool use human in the loop",
            "site:paperswithcode.com retrieval augmented generation evaluation",
        ],
    },
}

LOW_COST_POC_OPTIONAL_SERVICES = {
    "azure api management",
    "application gateway",
    "azure front door",
    "azure kubernetes service",
    "private endpoint",
    "virtual network",
    "private dns",
}


ARCHITECTURE_STACK = [
    "Gemini API or local LLM fallback for extraction and drafting",
    "FAISS, Chroma, or pgvector for retrieval over proposal documents",
    "Local disk or S3-compatible storage for uploads and generated PPTX files",
    "PostgreSQL for projects, reviews, and version history",
    "Optional Tavily or SerpAPI research tools when external evidence is needed",
    "Queue or background worker for long-running document jobs",
    "PowerPoint template renderer for branded enterprise decks",
    "Application telemetry and logging for quality and traceability",
]


def _gemini_available() -> bool:
    return bool(os.getenv("GEMINI_API_KEY"))


def _extract_json_object(text: str) -> Dict[str, Any]:
    cleaned = text.strip()
    if cleaned.startswith("```"):
        cleaned = re.sub(r"^```(?:json)?", "", cleaned, flags=re.IGNORECASE).strip()
        cleaned = re.sub(r"```$", "", cleaned).strip()
    try:
        return json.loads(cleaned)
    except json.JSONDecodeError:
        match = re.search(r"\{.*\}", cleaned, flags=re.DOTALL)
        if not match:
            raise
        return json.loads(match.group(0))


def _extract_with_gemini(document_text: str, prompt: str) -> Dict[str, Any]:
    model = os.getenv("GEMINI_MODEL", "gemini-2.5-flash")
    api_key = os.getenv("GEMINI_API_KEY")
    url = f"https://generativelanguage.googleapis.com/v1beta/models/{model}:generateContent?key={api_key}"
    payload = {
        "contents": [
            {
                "role": "user",
                "parts": [
                    {
                        "text": f"{prompt}\n\nCustomer document:\n{document_text[:45000]}",
                    }
                ],
            }
        ],
        "generationConfig": {
            "temperature": 0.1,
            "responseMimeType": "application/json",
        },
    }
    request = urllib.request.Request(
        url,
        data=json.dumps(payload).encode("utf-8"),
        headers={"Content-Type": "application/json"},
        method="POST",
    )
    try:
        with urllib.request.urlopen(request, timeout=60) as response:
            raw = json.loads(response.read().decode("utf-8"))
    except urllib.error.HTTPError as exc:
        detail = exc.read().decode("utf-8", errors="ignore")
        raise RuntimeError(f"Gemini API error {exc.code}: {detail}") from exc

    candidates = raw.get("candidates", [])
    if not candidates:
        raise RuntimeError("Gemini API returned no candidates.")
    parts = candidates[0].get("content", {}).get("parts", [])
    text = "".join(part.get("text", "") for part in parts)
    return _extract_json_object(text)


def _sentences_with_keywords(text: str, keywords: List[str]) -> List[str]:
    sentences = re.split(r"(?<=[.!?])\s+|\n+", text)
    matches = []
    for sentence in sentences:
        clean = sentence.strip(" -\t\r")
        if clean and any(keyword.lower() in clean.lower() for keyword in keywords):
            matches.append(clean)
    return matches[:8]


def fallback_extract_requirements(document_text: str) -> Dict[str, List[str]]:
    text = document_text.strip()
    extracted = {
        "business_goals": _sentences_with_keywords(
            text, ["goal", "objective", "purpose", "business", "improve", "reduce", "increase", "optimize"]
        ),
        "functional_requirements": _sentences_with_keywords(
            text, ["must", "shall", "can", "user", "upload", "answer", "generate", "integrate", "workflow"]
        ),
        "non_functional_requirements": _sentences_with_keywords(
            text, ["scale", "latency", "performance", "availability", "monitoring", "cost", "maintainability"]
        ),
        "constraints": _sentences_with_keywords(
            text, ["constraint", "budget", "timeline", "deadline", "region", "policy"]
        ),
        "assumptions": [
            "The document is readable in supported formats or can be converted from text extraction.",
            "Human review is required before final proposal export.",
            "The proposal can be generated from a single prompt pass plus review refinements.",
        ],
        "integrations": _sentences_with_keywords(
            text, ["integrate", "identity", "SSO", "API", "CRM", "ERP", "SharePoint", "Teams", "database"]
        ),
        "security_needs": _sentences_with_keywords(
            text, ["security", "RBAC", "role", "access", "encrypt", "audit", "private", "compliance"]
        ),
        "expected_users": _sentences_with_keywords(
            text, ["user", "employee", "admin", "architect", "thousand", "concurrent", "role"]
        ),
    }

    for field in REQUIREMENT_FIELDS:
        if not extracted[field]:
            extracted[field] = []
    return extracted


def extract_requirements(document_text: str) -> Dict[str, Any]:
    prompt = f"""
You are a requirement extraction agent. Extract only customer requirements that are
explicitly supported by the document. Return concise valid JSON only.

Classification is mutually exclusive: each distinct source obligation belongs in exactly
one primary category. Do not repeat it in another category.

Primary-category rules:
- business_goals: explicit organisational outcome or KPI only; not a requested capability.
- functional_requirements: a user-visible capability or workflow.
- non_functional_requirements: measurable quality, scale, latency, availability, retention,
  observability, or maintainability.
- constraints: deadline, data residency, mandated technology, policy, budget, or procurement limit.
- integrations: named external identity provider, repository, API, or enterprise system.
- security_needs: authentication, authorization, confidentiality, access control, or audit control.
- expected_users: named human persona and its stated access boundary only. Capacity/concurrency
  is not an expected user; classify it as non_functional_requirement.
- assumptions: only statements explicitly labelled as an assumption by the customer.

Deduplication rules:
- Merge identical or semantically equivalent sentences, including repeated wording in the input.
- Keep attributes together. For example, "3,000 concurrent users with 95% under 3 seconds"
  becomes two requirements only because capacity and latency are independently testable.
- Do not create restatements such as putting RBAC in functional_requirements, integrations,
  and security_needs. Put it in security_needs; put Microsoft Entra ID itself in integrations.

Never infer or add admin UI, encryption, MFA, SIEM, APIs/connectors, citations, roles, or
observability merely because they are common. Do not turn implementation suggestions into
customer requirements. Prefer fewer accurate items to a comprehensive-looking list.

Return valid JSON only with:
{{
  "extracted": {{
    "business_goals": [],
    "functional_requirements": [],
    "non_functional_requirements": [],
    "constraints": [],
    "assumptions": [],
    "integrations": [],
    "security_needs": [],
    "expected_users": []
  }},
  "quality_notes": {{"overall": "short assessment"}}
}}

Quality guide:
{json.dumps(FIELD_QUALITY_GUIDE, indent=2)}
"""

    if azure_openai_configured():
        try:
            parsed, _ = create_json_response(f"{prompt}\n\nCustomer document:\n{document_text[:45000]}")
            parsed["extracted"] = _sanitize_extracted_requirements(parsed.get("extracted", {}))
            return {"mode": "azure_openai", **parsed, "quality_guide": FIELD_QUALITY_GUIDE}
        except Exception as exc:
            return {
                "mode": "fallback",
                "provider_error": str(exc),
                "extracted": fallback_extract_requirements(document_text),
                "quality_guide": FIELD_QUALITY_GUIDE,
            }

    if _gemini_available():
        try:
            parsed = _extract_with_gemini(document_text, prompt)
            return {"mode": "gemini", **parsed, "quality_guide": FIELD_QUALITY_GUIDE}
        except Exception as exc:
            return {
                "mode": "fallback",
                "provider_error": str(exc),
                "extracted": fallback_extract_requirements(document_text),
                "quality_guide": FIELD_QUALITY_GUIDE,
            }

    return {
        "mode": "fallback",
        "extracted": fallback_extract_requirements(document_text),
        "quality_guide": FIELD_QUALITY_GUIDE,
    }


def extract_requirements_from_chunks(chunks: List[Dict[str, Any]]) -> Dict[str, Any]:
    """Extract requirements exclusively through the published Foundry Requirement Agent."""
    if not published_agent_configured("requirement"):
        return _published_requirement_failure("Published Requirement Agent is not configured.")

    try:
        clean_chunks = [chunk for chunk in chunks if str(chunk.get("text", "")).strip()]
        agent_outputs = [_invoke_requirement_agent_batch([chunk]) for chunk in clean_chunks] if len(clean_chunks) > 1 else [
            _invoke_requirement_agent_batch(clean_chunks)
        ]
        agent_output = _merge_requirement_agent_outputs(agent_outputs)
        raw_requirements = agent_output.get("requirements", {})
        extracted = _sanitize_extracted_requirements(_requirement_statements(raw_requirements))
        canonical_records = _canonical_agent_requirement_records(raw_requirements, extracted)
        questions = _completion_questions_from_agent(agent_output) + _global_conflict_questions(extracted)
        return {
            "mode": "foundry_published_requirement_agent",
            "extracted": extracted,
            # Keep the raw Foundry response in the caller's checkpoint if needed, but only
            # pass this canonical, de-duplicated record set to later workflow stages.
            "agent_requirements": canonical_records,
            "traceability": _traceability_from_agent_records(canonical_records, extracted, chunks),
            "clarification_questions": _dedupe_questions(questions)[:3],
            "batch_count": len(agent_outputs),
            "provider_errors": [],
            "quality_guide": FIELD_QUALITY_GUIDE,
            "quality_notes": agent_output.get("summary", agent_output.get("quality_notes", {})),
        }
    except Exception as exc:
        return _published_requirement_failure(f"Published Requirement Agent failed: {exc}")


def _invoke_requirement_agent_batch(chunks: List[Dict[str, Any]]) -> Dict[str, Any]:
    return invoke_published_agent(
        "requirement",
        {
            "document_context": {
                "input_type": "requirement_document",
                "data_sensitivity": "customer_provided",
                "extraction_scope": "whole_document" if len(chunks) == 1 and str(chunks[0].get("section")) == "Whole Document" else "chunk",
                "global_merge_follows": True,
            },
            "chunks": [
                {
                    "chunk_id": chunk.get("chunk_id"),
                    "section": chunk.get("section"),
                    "heading_path": chunk.get("heading_path"),
                    "page_start": chunk.get("page_start"),
                    "page_end": chunk.get("page_end"),
                    "estimated_tokens": chunk.get("estimated_tokens"),
                    "text": str(chunk.get("text", "")),
                }
                for chunk in chunks
                if str(chunk.get("text", "")).strip()
            ],
        },
    )


def _merge_requirement_agent_outputs(agent_outputs: List[Dict[str, Any]]) -> Dict[str, Any]:
    merged_requirements: Dict[str, List[Any]] = {field: [] for field in REQUIREMENT_FIELDS}
    questions: List[Any] = []
    summaries: List[Any] = []
    for output in agent_outputs:
        requirements = output.get("requirements") or {}
        if isinstance(requirements, dict):
            for field in REQUIREMENT_FIELDS:
                values = requirements.get(field) or []
                if isinstance(values, list):
                    merged_requirements[field].extend(values)
        form = output.get("completion_form") or output.get("questions") or []
        if isinstance(form, list):
            questions.extend(form)
        summary = output.get("summary") or output.get("quality_notes")
        if summary:
            summaries.append(summary)
    return {
        "requirements": merged_requirements,
        "completion_form": questions,
        "summary": {"chunk_summaries": summaries} if len(summaries) > 1 else (summaries[0] if summaries else {}),
    }


def _published_requirement_failure(error: str) -> Dict[str, Any]:
    return {
        "mode": "failed",
        "error": error,
        "extracted": {field: [] for field in REQUIREMENT_FIELDS},
        "traceability": {},
        "batch_count": 0,
        "provider_errors": [error],
        "quality_guide": FIELD_QUALITY_GUIDE,
    }


def _requirement_statements(requirements: Any) -> Dict[str, List[str]]:
    """Convert Foundry Agent records into the compact UI format without losing raw records."""
    statements: Dict[str, List[str]] = {field: [] for field in REQUIREMENT_FIELDS}
    if not isinstance(requirements, dict):
        return statements
    for field in REQUIREMENT_FIELDS:
        for item in requirements.get(field, []):
            if isinstance(item, dict):
                statement = str(item.get("statement", "")).strip()
            else:
                statement = str(item).strip()
            if statement:
                statements[field].append(statement)
    return statements


def _requirements_with_ids(extracted: Dict[str, List[str]]) -> Dict[str, List[Dict[str, str]]]:
    """Provide a deterministic fallback contract when a legacy extraction is used."""
    records: Dict[str, List[Dict[str, str]]] = {field: [] for field in REQUIREMENT_FIELDS}
    index = 1
    for field in REQUIREMENT_FIELDS:
        for statement in extracted.get(field, []):
            records[field].append({"id": f"REQ-{index:03d}", "statement": str(statement)})
            index += 1
    return records


def _canonical_agent_requirement_records(
    requirements: Any, extracted: Dict[str, List[str]]
) -> Dict[str, List[Dict[str, Any]]]:
    """Retain the source agent's IDs while aligning records with the canonical UI set.

    A Foundry agent can put the same obligation in two categories (for example,
    SharePoint under both functional requirements and integrations).  The compact
    ``extracted`` set already resolves that ambiguity.  This function makes the
    record-level contract follow the same decision before Architect receives it.
    """
    canonical: Dict[str, List[Dict[str, Any]]] = {field: [] for field in REQUIREMENT_FIELDS}
    candidates: List[Dict[str, Any]] = []
    if isinstance(requirements, dict):
        for proposed_field in REQUIREMENT_FIELDS:
            for item in requirements.get(proposed_field, []):
                if not isinstance(item, dict):
                    continue
                statement = str(item.get("statement", "")).strip()
                if not statement:
                    continue
                record = dict(item)
                record["statement"] = statement
                record["_canonical_category"] = _canonical_requirement_category(statement, proposed_field)
                candidates.append(record)

    used_positions: set[int] = set()
    generated_id = 1
    for field in REQUIREMENT_FIELDS:
        for statement in extracted.get(field, []):
            match_index = next(
                (
                    index
                    for index, item in enumerate(candidates)
                    if index not in used_positions
                    and item["_canonical_category"] == field
                    and _same_requirement(str(item["statement"]), statement)
                ),
                None,
            )
            if match_index is None:
                match_index = next(
                    (
                        index
                        for index, item in enumerate(candidates)
                        if index not in used_positions
                        and _same_requirement(str(item["statement"]), statement)
                    ),
                    None,
                )
            if match_index is None:
                record: Dict[str, Any] = {
                    "id": f"REQ-CANON-{generated_id:03d}",
                    "statement": statement,
                    "source_chunk_ids": [],
                    "confidence": "medium",
                }
                generated_id += 1
            else:
                used_positions.add(match_index)
                record = dict(candidates[match_index])
                record.pop("_canonical_category", None)
                record["statement"] = statement
            canonical[field].append(record)
    return canonical


def _traceability_from_agent_records(
    requirements: Any, extracted: Dict[str, List[str]], chunks: List[Dict[str, Any]]
) -> Dict[str, List[Dict[str, Any]]]:
    """Use the agent's chunk IDs when present, then fill any gaps deterministically."""
    fallback = _build_traceability(extracted, chunks)
    if not isinstance(requirements, dict):
        return fallback

    records: Dict[str, List[Dict[str, Any]]] = {field: [] for field in REQUIREMENT_FIELDS}
    for field in REQUIREMENT_FIELDS:
        source_items = [item for item in requirements.get(field, []) if isinstance(item, dict)]
        for position, statement in enumerate(extracted.get(field, [])):
            source_item = next(
                (item for item in source_items if str(item.get("statement", "")).strip() == statement),
                None,
            )
            if source_item:
                source = source_item.get("source", {}) if isinstance(source_item.get("source"), dict) else {}
                source_chunk_ids = source_item.get("source_chunk_ids") or (
                    [source.get("chunk_id")] if source.get("chunk_id") else []
                )
                records[field].append(
                    {
                        "id": source_item.get("id"),
                        "statement": statement,
                        "source_chunk_ids": [str(item) for item in source_chunk_ids if item],
                        "confidence": str(source_item.get("confidence") or "medium"),
                    }
                )
            else:
                records[field].append(fallback[field][position])
    return records


def _completion_questions_from_agent(agent_output: Dict[str, Any]) -> List[Dict[str, str]]:
    """Adapt the published Requirement Agent's completion form to the existing HITL UI."""
    questions = agent_output.get("completion_form") or agent_output.get("questions") or []
    normalized = []
    for item in questions:
        if not isinstance(item, dict) or not item.get("question"):
            continue
        normalized.append(
            {
                "id": str(item.get("id") or f"Q-{len(normalized) + 1:03d}"),
                "requirement_id": str(item.get("requirement_id") or ""),
                "field": str(item.get("field") or "clarification"),
                "question": str(item["question"]),
                "reason": str(item.get("impact") or item.get("why_it_matters") or "Needed before architecture approval."),
                "answer_type": str(item.get("answer_type") or "text"),
                "options": [str(option) for option in item.get("options", item.get("suggested_options", [])) if str(option).strip()],
            }
        )
    return normalized[:3]


def _global_conflict_questions(extracted: Dict[str, List[str]]) -> List[Dict[str, str]]:
    all_statements = [
        statement
        for values in extracted.values()
        for statement in values
        if isinstance(statement, str)
    ]
    joined = " ".join(all_statements).lower()
    questions = []
    if re.search(r"\b30\s+days?\b", joined) and re.search(r"\b90\s+days?\b", joined):
        questions.append(
            {
                "id": "GLOBAL-RETENTION-001",
                "requirement_id": "",
                "field": "retention_policy",
                "question": "Which retention policy is authoritative for uploaded customer source documents: 30 days, 90 days, or another period?",
                "reason": "Conflicting retention requirements were found across document sections.",
                "answer_type": "single_select",
                "options": ["30 days", "90 days", "Other"],
            }
        )
    if "availability" in joined and any(term in joined for term in ["to be confirmed", "tbd", "confirm during solution design"]):
        questions.append(
            {
                "id": "GLOBAL-AVAILABILITY-001",
                "requirement_id": "",
                "field": "availability_target",
                "question": "What availability target should the POC architecture use until production SLA is finalized?",
                "reason": "Availability is referenced but left open for solution design.",
                "answer_type": "text",
                "options": [],
            }
        )
    return questions


def _dedupe_questions(questions: List[Dict[str, str]]) -> List[Dict[str, str]]:
    deduped = []
    seen = set()
    for question in questions:
        key = _requirement_fingerprint(str(question.get("field") or "") + " " + str(question.get("question") or ""))
        if not key or key in seen:
            continue
        seen.add(key)
        deduped.append(question)
    return deduped


def _build_traceability(extracted: Dict[str, List[str]], chunks: List[Dict[str, Any]]) -> Dict[str, List[Dict[str, Any]]]:
    traceability: Dict[str, List[Dict[str, Any]]] = {field: [] for field in REQUIREMENT_FIELDS}
    for field in REQUIREMENT_FIELDS:
        for statement in extracted.get(field, []):
            tokens = {
                token
                for token in re.findall(r"[a-z0-9]{4,}", statement.lower())
                if token not in {"must", "shall", "with", "from", "that", "this", "system"}
            }
            scored = []
            for chunk in chunks:
                haystack = str(chunk.get("text", "")).lower()
                score = sum(token in haystack for token in tokens)
                if score:
                    scored.append((score, chunk))
            source_chunks = [
                item[1]["chunk_id"] for item in sorted(scored, key=lambda item: item[0], reverse=True)[:2]
            ]
            traceability[field].append(
                {"statement": statement, "source_chunk_ids": source_chunks, "confidence": "high" if source_chunks else "medium"}
            )
    return traceability


def build_clarification_questions(extracted: Dict[str, List[str]]) -> List[Dict[str, str]]:
    questions = []
    checks = [
        ("business_goals", "What business outcomes or KPIs should this proposal optimize for?"),
        ("expected_users", "How many users, roles, and concurrent sessions should the solution support?"),
        ("security_needs", "What identity provider, RBAC model, audit, and data isolation requirements are mandatory?"),
        ("integrations", "Which enterprise systems or document repositories must be integrated first?"),
        ("non_functional_requirements", "What are the target availability, latency, retention, and monitoring requirements?"),
        ("constraints", "Are there fixed budget, timeline, region, data residency, or technology constraints?"),
    ]
    for field, question in checks:
        values = extracted.get(field, [])
        if not values:
            questions.append(
                {
                    "field": field,
                    "question": question,
                    "reason": FIELD_QUALITY_GUIDE[field],
                }
            )
    return questions


def recommend_patterns(extracted: Dict[str, List[str]]) -> List[Dict[str, str]]:
    text = " ".join(" ".join(v) for v in extracted.values()).lower()
    recommendations = []
    for pattern in AI_PATTERN_KB:
        name = pattern["name"].lower()
        include = (
            ("document" in text and "rag" in name)
            or ("role" in text and "rag" in name)
            or ("workflow" in text and "workflow" in name)
            or ("agent" in text and "agent" in name)
            or ("legacy" in text and "computer-use" in name)
            or ("relationship" in text and "graph" in name)
        )
        if include:
            recommendations.append(pattern)
    return recommendations or [AI_PATTERN_KB[0], AI_PATTERN_KB[4]]


def _fallback_slide_plan(proposal: Dict[str, Any]) -> List[Dict[str, Any]]:
    slides = []
    for section_key, title in PROPOSAL_SECTIONS:
        value = proposal.get(section_key, "")
        if not value:
            continue
        bullets = [str(item) for item in value[:5]] if isinstance(value, list) else [str(value)]
        slides.append({"section_key": section_key, "title": title, "bullets": bullets, "requirement_ids": [], "source_urls": []})
    return slides


def fallback_generate_proposal(
    extracted: Dict[str, List[str]],
    clarifications: Optional[Dict[str, str]] = None,
    architecture_decisions: Optional[Dict[str, Any]] = None,
    cost_assumptions: Optional[List[Dict[str, Any]]] = None,
    template_id: str = "starter",
    architecture_approved: bool = False,
) -> Dict[str, Any]:
    clarifications = clarifications or {}
    patterns = recommend_patterns(extracted)
    expected_output = [
        "Extract the key requirements from the uploaded document.",
        "Ask clarification questions when security, scale, user roles, data source, or integration details are missing.",
        "Propose an implementable AI architecture and deployment design.",
        "Estimate costs for Development, Test, and UAT or Staging environments.",
        "Export a customer-ready PowerPoint proposal using the selected enterprise template while preserving branding.",
    ]
    key_engineering_challenges = [
        "Understanding unstructured requirement documents accurately.",
        "Avoiding generic, shallow, or hallucinated solution recommendations.",
        "Keeping the proposal consistent across business solution, technical architecture, infrastructure, and cost estimation.",
        "Mapping AI-generated content into PowerPoint templates without breaking the design.",
        "Supporting multiple proposal templates and future template changes.",
        "Creating a useful HITL loop so users can refine assumptions, solution direction, architecture, and final slide content.",
    ]

    proposal = {
        "proposal_id": str(uuid.uuid4()),
        "generated_at": datetime.utcnow().isoformat() + "Z",
        "template_id": template_id,
        "architecture_approved": architecture_approved,
        "architecture_decisions": architecture_decisions or {},
        "cost_assumptions": cost_assumptions or [],
        "executive_summary": (
            "Build a secure AI Solution Architect proposal-generation agent that ingests customer requirement "
            "documents, extracts structured requirements, asks targeted clarification questions, recommends an "
            "implementable solution architecture, and exports a customer-ready PowerPoint using enterprise templates."
        ),
        "problem_statement": (
            "Proposal teams need to convert long, unstructured requirement documents into consistent business, "
            "technical, architecture, cost, risk, and assumption content without breaking approved presentation branding."
        ),
        "expected_output": expected_output,
        "key_engineering_challenges": key_engineering_challenges,
        "business_solution": [
            "Document intake with OCR/layout parsing and hierarchical chunking for long documents.",
            "Structured requirement extraction against a field-quality knowledge base.",
            "Human-in-the-loop clarification before final architecture and proposal writing.",
            "Reviewable proposal sections before PPTX export.",
            "Template abstraction layer for future slide templates.",
        ],
        "technical_solution": [
            "Ingestion service parses Word, PDF, text, and scanned content into normalized document sections.",
            "Extraction agent produces JSON for business goals, functional and non-functional requirements, constraints, assumptions, integrations, security needs, and expected users.",
            "Solution architect agent selects AI patterns from a controlled pattern knowledge base.",
            "Proposal composer generates section-specific content and validates consistency across architecture, cost, risks, and assumptions.",
            "PPT renderer applies template mappings and only updates intended placeholders.",
        ],
        "recommended_patterns": patterns,
        "architecture_stack": ARCHITECTURE_STACK,
        "infrastructure_design": [
            "Containerized API hosts the proposal workflow and worker tasks.",
            "A vector database stores document chunks and metadata for retrieval.",
            "Object storage keeps source documents, OCR artifacts, and generated proposal files.",
            "PostgreSQL stores project metadata, version history, and review state.",
            "A secrets manager protects API keys and connection strings.",
            "A queue or job runner coordinates parsing, extraction, review, and PPT export jobs.",
            "Application telemetry tracks latency, token usage, failures, and proposal-generation quality signals.",
        ],
        "cost_estimation": cost_assumptions or [
            {"environment": "Development", "monthly_estimate_usd": "TBD", "notes": "Provide approved capacity and usage assumptions before estimating cost."},
            {"environment": "Test", "monthly_estimate_usd": "TBD", "notes": "Provide approved capacity and usage assumptions before estimating cost."},
            {"environment": "UAT/Staging", "monthly_estimate_usd": "TBD", "notes": "Provide approved capacity and usage assumptions before estimating cost."},
        ],
        "implementation_approach": [
            "Phase 1: Template inventory, schema mapping, document ingestion MVP.",
            "Phase 2: Extraction, HITL clarification, solution-pattern recommendation.",
            "Phase 3: Proposal composer, consistency checks, cost model, PPTX export.",
            "Phase 4: Hardening, monitoring, UAT, and rollout.",
        ],
        "risks": [
            "Poor source document quality may reduce extraction accuracy; mitigate with OCR/layout confidence checks and reviewer approval.",
            "Template drift may break mappings; mitigate with template schema validation and visual regression checks.",
            "Generic architecture output may reduce trust; mitigate with controlled pattern KB and required assumptions.",
            "Cost estimates can vary with token volume and document count; mitigate with usage telemetry and environment budgets.",
        ],
        "assumptions": extracted.get("assumptions", []) + list(clarifications.values()),
        "delivery_notes": [
            "Use a branded enterprise template when available; fall back to the starter deck when no template is configured.",
            "Keep slide copy concise and audience-facing so the output reads like a real proposal.",
            "Prefer cloud-neutral wording unless a customer explicitly asks for a vendor-specific deployment.",
        ],
        "extracted_requirements": extracted,
    }
    proposal["slide_plan"] = build_slide_plan(proposal)
    return proposal


def _sanitize_extracted_requirements(extracted: Any) -> Dict[str, List[str]]:
    """Normalize category boundaries and remove duplicate obligations across LLM batches."""
    cleaned: Dict[str, List[str]] = {field: [] for field in REQUIREMENT_FIELDS}
    for field in REQUIREMENT_FIELDS:
        values = extracted.get(field, []) if isinstance(extracted, dict) else []
        if not isinstance(values, list):
            continue
        for value in values:
            text = str(value).strip()
            if not text:
                continue
            category = _canonical_requirement_category(text, field)
            if category == "expected_users" and any(
                term in text.lower() for term in ("concurrent", "capacity", "throughput")
            ):
                continue
            if category == "assumptions" and not re.search(
                r"\b(assum|expected to|likely|provisional|pending confirmation)\b", text, flags=re.IGNORECASE
            ):
                continue
            if any(_same_requirement(text, existing) for existing in cleaned[category]):
                continue
            for other_field in REQUIREMENT_FIELDS:
                if other_field == category:
                    continue
                if any(_same_requirement(text, existing) for existing in cleaned[other_field]):
                    break
            else:
                cleaned[category].append(text)
    return cleaned


def _canonical_requirement_category(text: str, proposed: str) -> str:
    value = text.lower()
    if any(marker in value for marker in ("confidential", "authorized", "rbac", "single sign-on", " sso", "audit", "encrypt", "public web search")):
        return "security_needs"
    if any(marker in value for marker in ("sharepoint", "salesforce", "microsoft graph", "integrate with", "external api")):
        return "integrations"
    if any(marker in value for marker in ("p95", "concurrent", "latency", "response time", "scalable", "user-friendly", "availability", "process a ")):
        return "non_functional_requirements"
    if any(marker in value for marker in ("azure as", "primary cloud", "monthly operating cost", "budget", "retained for", "deleted after", "data residency", "deadline")):
        return "constraints"
    if re.search(r"\b(assum|expected to|likely|provisional|pending confirmation)\b", value):
        return "assumptions"
    return proposed if proposed in REQUIREMENT_FIELDS else "functional_requirements"


def _same_requirement(first: str, second: str) -> bool:
    first_normalized = _requirement_fingerprint(first)
    second_normalized = _requirement_fingerprint(second)
    if first_normalized == second_normalized:
        return True
    if "architecture generation" in first.lower() and "architecture generation" in second.lower():
        return "approv" in first.lower() and "approv" in second.lower()
    first_terms = set(first_normalized.split())
    second_terms = set(second_normalized.split())
    if not first_terms or not second_terms:
        return False
    return len(first_terms & second_terms) / min(len(first_terms), len(second_terms)) >= 0.86


def _requirement_fingerprint(text: str) -> str:
    stop_words = {"the", "system", "platform", "solution", "must", "shall", "should", "also", "allow", "to", "and", "or", "a", "an"}
    tokens = [
        token for token in re.findall(r"[a-z0-9]+", text.lower()) if token not in stop_words
    ]
    return " ".join(tokens)


def generate_proposal(
    extracted: Dict[str, List[str]],
    clarifications: Optional[Dict[str, str]] = None,
    architecture_decisions: Optional[Dict[str, Any]] = None,
    cost_assumptions: Optional[List[Dict[str, Any]]] = None,
    template_id: str = "starter",
    architecture_approved: bool = False,
    confirmed_requirements: Optional[Dict[str, Any]] = None,
) -> Dict[str, Any]:
    """Draft proposal content with the Published Proposal Agent and retain renderer-compatible metadata."""
    fallback = fallback_generate_proposal(
        extracted,
        clarifications,
        architecture_decisions,
        cost_assumptions,
        template_id,
        architecture_approved,
    )
    if published_agent_configured("proposal"):
        try:
            raw_proposal = invoke_published_agent(
                "proposal",
                {
                    "customer_context": {"input_type": "approved_customer_requirement_document"},
                    "confirmed_requirements": confirmed_requirements or _requirements_with_ids(extracted),
                    "confirmed_hitl_answers": clarifications or {},
                    "approved_architecture": architecture_decisions or {},
                    "cost_assumptions": cost_assumptions or [],
                    "template_rules": {
                        "template_id": template_id,
                        "max_bullets_per_slide": 5,
                        "component_card_max_lines": 3,
                        "preserve_branding": True,
                        "preserve_placeholder_positions": True,
                        "supported_layout_ids": [
                            "title", "customer_context", "problem_impact", "goals_kpi",
                            "four_column_value_chain", "process_flow", "logical_architecture",
                            "component_cards", "infrastructure_security", "cost_table", "roadmap",
                        ],
                    },
                },
            )
            pptagent_input = raw_proposal.get("pptagent_input", raw_proposal)
            if not isinstance(pptagent_input, dict) or not isinstance(pptagent_input.get("slides"), list):
                raise ValueError("Published Proposal Agent returned no pptagent_input.slides array.")
            fallback["pptagent_input"] = pptagent_input
            fallback["proposal_agent_output"] = raw_proposal
            fallback["generation_mode"] = "foundry_published_proposal_agent"
            return fallback
        except Exception as exc:
            fallback["generation_mode"] = "failed"
            fallback["generation_error"] = str(exc)
            return fallback
    if not azure_openai_configured():
        fallback["generation_mode"] = "fallback"
        return fallback

    prompt = f"""
You are an enterprise proposal-writing agent. Create concise, customer-ready proposal
content from approved requirements, approved clarification answers, and approved
research evidence. Do not use web search. Do not invent prices, citations, or facts.

Requirements JSON:
{json.dumps(extracted, indent=2)}

Clarification answers JSON:
{json.dumps(clarifications or {}, indent=2)}

Approved architecture decisions and research evidence JSON:
{json.dumps(architecture_decisions or {}, indent=2)}

Approved cost assumptions JSON:
{json.dumps(cost_assumptions or [], indent=2)}

Selected template ID:
{template_id}

Return JSON only with these keys:
{{
  "executive_summary": "string",
  "problem_statement": "string",
  "business_solution": ["string"],
  "technical_solution": ["string"],
  "infrastructure_design": ["string"],
  "cost_estimation": [{{"environment": "string", "monthly_estimate_usd": "string", "notes": "string"}}],
  "implementation_approach": ["string"],
  "risks": ["string"],
  "delivery_notes": ["string"],
  "slide_plan": [{{
    "section_key": "executive_summary|problem_statement|business_solution|technical_solution|infrastructure_design|cost_estimation|implementation_approach|risks",
    "title": "string",
    "bullets": ["string"],
    "requirement_ids": ["string"],
    "source_urls": ["string"]
  }}]
}}

Keep bullets specific, concise, and consistent with the approved evidence. Flag an
assumption rather than presenting it as a confirmed customer fact. Do not invent costs;
use only approved cost assumptions or return TBD. Keep at most five bullets per slide.
"""
    try:
        generated, _ = create_json_response(prompt)
    except Exception:
        fallback["generation_mode"] = "fallback"
        return fallback

    allowed_keys = {
        "executive_summary",
        "problem_statement",
        "business_solution",
        "technical_solution",
        "infrastructure_design",
        "cost_estimation",
        "implementation_approach",
        "risks",
        "delivery_notes",
        "slide_plan",
    }
    for key in allowed_keys:
        if key in generated and generated[key]:
            fallback[key] = generated[key]
    # The LLM drafts proposal copy; deterministic planning owns narrative order,
    # layout selection, and compact renderer-safe slide content.
    fallback["slide_plan"] = build_slide_plan(fallback)
    fallback["generation_mode"] = "azure_openai"
    return fallback


def _normalize_slide_plan(slide_plan: Any, proposal: Dict[str, Any]) -> List[Dict[str, Any]]:
    allowed_sections = {key for key, _ in PROPOSAL_SECTIONS}
    valid_slides = []
    for slide in slide_plan or []:
        if not isinstance(slide, dict) or slide.get("section_key") not in allowed_sections:
            continue
        bullets = [str(item).strip() for item in slide.get("bullets", []) if str(item).strip()][:5]
        valid_slides.append({
            "section_key": slide["section_key"],
            "title": str(slide.get("title") or slide["section_key"].replace("_", " ").title()),
            "bullets": bullets,
            "requirement_ids": [str(item) for item in slide.get("requirement_ids", [])],
            "source_urls": [str(item) for item in slide.get("source_urls", [])],
        })
    return valid_slides or _fallback_slide_plan(proposal)


def _active_requirement_fields(extracted: Dict[str, List[str]]) -> List[str]:
    return [field for field in REQUIREMENT_FIELDS if extracted.get(field)]


def _infer_architecture_patterns(extracted: Dict[str, List[str]]) -> List[str]:
    text = " ".join(" ".join(values) for values in extracted.values()).lower()
    patterns: List[str] = []
    if any(term in text for term in ("document", "citation", "knowledge", "answer questions", "search")):
        patterns.append("RAG with metadata filtering")
    if any(term in text for term in ("approval", "review", "workflow", "proposal", "upload")):
        patterns.append("Workflow automation")
    if any(term in text for term in ("pdf", "ocr", "scanned", "table", "long-running", "150-page")):
        patterns.append("Asynchronous document processing")
    return patterns or ["Workflow automation"]


def _is_official_azure_source(url: str) -> bool:
    host = url.lower().split("//", 1)[-1].split("/", 1)[0]
    if any(host == source_host or host.endswith(f".{source_host}") for source_host in ("learn.microsoft.com", "azure.microsoft.com")):
        return True
    lowered = url.lower()
    return lowered.startswith("https://github.com/azure/") or lowered.startswith("https://github.com/microsoft/")


def _source_host(url: str) -> str:
    return url.lower().split("//", 1)[-1].split("/", 1)[0]


def _classify_source(url: str) -> str:
    host = _source_host(url)
    if host in {"learn.microsoft.com", "azure.microsoft.com"} or host.endswith(".microsoft.com"):
        return "azure_docs"
    if host == "github.com":
        lowered = url.lower()
        if lowered.startswith("https://github.com/azure/") or lowered.startswith("https://github.com/microsoft/"):
            return "microsoft_github"
        return "github"
    if host in {"arxiv.org", "aclanthology.org", "openreview.net", "paperswithcode.com"}:
        return "research_paper"
    return "web"


def _source_allowed(url: str, research_mode: str) -> bool:
    if research_mode == "off":
        return False
    if research_mode == "azure_official_only":
        return _is_official_azure_source(url)
    if research_mode == "azure_github_papers":
        source_type = _classify_source(url)
        return source_type in {"azure_docs", "microsoft_github", "research_paper"}
    return _is_official_azure_source(url)


def _normalize_research_mode(research_mode: str) -> str:
    return research_mode if research_mode in RESEARCH_MODES else "azure_official_only"


def _research_policy(research_mode: str) -> Dict[str, Any]:
    normalized_mode = _normalize_research_mode(research_mode)
    policy = dict(ARCHITECT_RESEARCH_POLICIES[normalized_mode])
    policy["never_send_customer_content_to_web_search"] = True
    return policy


def _normalize_sources(sources: Any, research_mode: str) -> List[Dict[str, str]]:
    normalized: List[Dict[str, str]] = []
    seen = set()
    for source in sources or []:
        if isinstance(source, str):
            source = {"url": source, "title": source}
        if not isinstance(source, dict):
            continue
        url = str(source.get("url", "")).strip()
        if not url or url in seen or not _source_allowed(url, research_mode):
            continue
        source_type = str(source.get("source_type") or _classify_source(url))
        normalized.append(
            {
                "title": str(source.get("title") or url),
                "url": url,
                "source_type": source_type,
                "summary": str(source.get("summary") or "Source used to validate architecture guidance."),
            }
        )
        seen.add(url)
    return normalized


def _default_cost_estimate() -> Dict[str, List[Dict[str, str]]]:
    note = "TBD until capacity and Azure Retail Prices API evidence are approved. Model token usage is reported separately."
    return {
        "development": [{"service": "POC platform services", "monthly_estimate_usd": "TBD", "notes": note}],
        "test": [{"service": "POC platform services", "monthly_estimate_usd": "TBD", "notes": note}],
        "uat_or_staging": [{"service": "POC platform services", "monthly_estimate_usd": "TBD", "notes": note}],
    }


def _normalize_cost_estimate(value: Any) -> Dict[str, List[Dict[str, str]]]:
    normalized = _default_cost_estimate()
    if not isinstance(value, dict):
        return normalized
    for environment in normalized:
        items = value.get(environment)
        if not isinstance(items, list) or not items:
            continue
        cleaned = []
        for item in items:
            if not isinstance(item, dict):
                continue
            cleaned.append(
                {
                    "service": str(item.get("service") or "Platform service"),
                    "monthly_estimate_usd": str(item.get("monthly_estimate_usd") or "TBD"),
                    "notes": str(item.get("notes") or "Verify against Azure Retail Prices API before approval."),
                }
            )
        if cleaned:
            normalized[environment] = cleaned
    return normalized


def _minimum_components(extracted: Dict[str, List[str]]) -> List[Dict[str, Any]]:
    active = _active_requirement_fields(extracted)
    components = [
        {
            "id": "ARC-API",
            "service": "Azure Container Apps",
            "classification": "required_for_poc",
            "purpose": "Host the authenticated API and deterministic proposal workflow.",
            "requirement_ids": active,
            "alternatives": ["Azure Functions for lightweight event handlers"],
        },
        {
            "id": "ARC-STORAGE",
            "service": "Azure Blob Storage",
            "classification": "required_for_poc",
            "purpose": "Store uploaded documents, parsed artifacts, and generated PowerPoint files.",
            "requirement_ids": [field for field in active if field in {"functional_requirements", "integrations", "constraints"}],
            "alternatives": ["Local disk for a local-only demo"],
        },
        {
            "id": "ARC-LLM",
            "service": "Microsoft Foundry / Azure OpenAI",
            "classification": "required_for_poc",
            "purpose": "Perform requirement extraction, architecture reasoning, and proposal drafting.",
            "requirement_ids": [field for field in active if field in {"functional_requirements", "business_goals"}],
            "alternatives": ["Provider abstraction for approved alternative models"],
        },
        {
            "id": "ARC-SECRETS",
            "service": "Azure Key Vault",
            "classification": "required_for_poc",
            "purpose": "Keep service credentials outside application code.",
            "requirement_ids": [field for field in active if field in {"security_needs", "constraints"}],
            "alternatives": ["Environment variables for local-only development"],
        },
        {
            "id": "ARC-MONITORING",
            "service": "Application Insights",
            "classification": "required_for_poc",
            "purpose": "Capture request latency, failures, workflow traces, and model usage.",
            "requirement_ids": [field for field in active if field in {"non_functional_requirements", "security_needs"}],
            "alternatives": ["Structured logs for a local-only demo"],
        },
        {
            "id": "ARC-TEMPLATE-MAPPING",
            "service": "Template mapping layer",
            "classification": "required_for_poc",
            "purpose": "Map approved proposal sections into predefined PowerPoint placeholders without changing template branding or layout.",
            "requirement_ids": ["functional_requirements"],
            "alternatives": ["Per-template JSON configuration"],
        },
    ]
    if extracted.get("integrations"):
        components.append(
            {
                "id": "ARC-INTEGRATIONS",
                "service": "Microsoft Graph / enterprise integration adapters",
                "classification": "required_for_poc",
                "purpose": "Connect approved enterprise document sources and destinations.",
                "requirement_ids": ["integrations"],
                "alternatives": ["Mock connector for a local-only demo"],
            }
        )
    if extracted.get("non_functional_requirements") or extracted.get("functional_requirements"):
        components.append(
            {
                "id": "ARC-ASYNC-WORKER",
                "service": "Azure Storage Queue and Container Apps job",
                "classification": "optional_for_poc",
                "purpose": "Run document parsing and PPTX generation asynchronously with retry support.",
                "requirement_ids": ["non_functional_requirements", "functional_requirements"],
                "alternatives": ["Azure Functions for lightweight jobs"],
            }
        )
    return components


def _canonical_service_name(service: str) -> str:
    lowered = service.lower()
    if "cognitive search" in lowered or "azure ai search" in lowered:
        return "Azure AI Search"
    if "azure ad" in lowered or "entra id" in lowered:
        return "Microsoft Entra ID"
    if "azure ai studio" in lowered or "azure openai" in lowered or "foundry" in lowered:
        return "Microsoft Foundry / Azure OpenAI"
    if "application insights" in lowered:
        return "Application Insights"
    if "azure key vault" in lowered:
        return "Azure Key Vault"
    return service.strip()


def _normalize_components(value: Any, extracted: Dict[str, List[str]]) -> List[Dict[str, Any]]:
    normalized = []
    seen_by_service: Dict[str, Dict[str, Any]] = {}
    for index, component in enumerate(value or [], start=1):
        if not isinstance(component, dict):
            continue
        service = str(component.get("service") or component.get("service_name") or "").strip()
        purpose = str(component.get("purpose") or "").strip()
        if not service or not purpose:
            continue
        classification = str(component.get("classification") or "optional_for_poc")
        service = _canonical_service_name(service)
        if service.lower() in LOW_COST_POC_OPTIONAL_SERVICES:
            classification = "optional_for_poc"
        alternatives = [str(item) for item in component.get("alternatives", []) if str(item).strip()]
        alternative = str(component.get("alternative") or "").strip()
        if alternative and alternative not in alternatives:
            alternatives.append(alternative)
        record = {
            "id": str(component.get("id") or f"ARC-{index:03d}"),
            "service": service,
            "classification": classification,
            "purpose": purpose,
            "requirement_ids": [str(item) for item in component.get("requirement_ids", component.get("mapped_requirements", []))],
            "alternatives": alternatives,
        }
        existing = seen_by_service.get(service.lower())
        if existing:
            existing["requirement_ids"] = list(dict.fromkeys(existing["requirement_ids"] + record["requirement_ids"]))
            existing["alternatives"] = list(dict.fromkeys(existing["alternatives"] + record["alternatives"]))
            if purpose not in existing["purpose"]:
                existing["purpose"] = f"{existing['purpose']} {purpose}"
            continue
        normalized.append(record)
        seen_by_service[service.lower()] = record
    return normalized or _minimum_components(extracted)


def _ensure_architecture_coverage(
    components: List[Dict[str, Any]], extracted: Dict[str, List[str]]
) -> List[Dict[str, Any]]:
    active_fields = set(_active_requirement_fields(extracted))
    covered_fields = {
        requirement_id
        for component in components
        for requirement_id in component.get("requirement_ids", [])
        if requirement_id in active_fields
    }
    component_text = " ".join(
        f"{component.get('service', '')} {component.get('purpose', '')}".lower()
        for component in components
    )
    semantic_coverage = {
        "business_goals": any(term in component_text for term in ("foundry", "azure openai", "proposal")),
        "functional_requirements": any(term in component_text for term in ("azure functions", "container apps", "logic apps", "workflow")),
        "non_functional_requirements": "application insights" in component_text or "monitor" in component_text,
        "constraints": "cost management" in component_text or "budget" in component_text,
        "assumptions": any(term in component_text for term in ("azure functions", "container apps", "workflow")),
        "integrations": any(term in component_text for term in ("microsoft graph", "integration", "connector")),
        "security_needs": "key vault" in component_text or "entra" in component_text,
        "expected_users": any(term in component_text for term in ("azure functions", "container apps", "api")),
    }
    covered_fields.update(field for field, covered in semantic_coverage.items() if covered)
    targets = {
        "business_goals": ("ARC-LLM", "Microsoft Foundry / Azure OpenAI", "Generate grounded solution content aligned to approved business outcomes."),
        "functional_requirements": ("ARC-API", "Azure Container Apps", "Host the authenticated API and deterministic workflow endpoints."),
        "non_functional_requirements": ("ARC-MONITORING", "Application Insights", "Measure latency, reliability, workload processing, and acceptance targets."),
        "constraints": ("ARC-COST-CONTROL", "Azure Cost Management", "Apply budget alerts and document cost assumptions for the POC."),
        "assumptions": ("ARC-API", "Azure Container Apps", "Record and expose approved delivery assumptions in the workflow."),
        "integrations": ("ARC-INTEGRATIONS", "Microsoft Graph / enterprise integration adapters", "Connect approved enterprise systems and data sources."),
        "security_needs": ("ARC-SECRETS", "Azure Key Vault and Microsoft Entra ID", "Protect credentials and enforce backend authorization."),
        "expected_users": ("ARC-API", "Azure Container Apps", "Serve authenticated users through the application API."),
    }
    for field in active_fields - covered_fields:
        target_id, service, purpose = targets[field]
        target = next((component for component in components if component["id"] == target_id), None)
        if target:
            target["requirement_ids"].append(field)
            continue
        components.append(
            {
                "id": target_id,
                "service": service,
                "classification": "required_for_poc",
                "purpose": purpose,
                "requirement_ids": [field],
                "alternatives": [],
            }
        )

    text = " ".join(" ".join(values) for values in extracted.values()).lower()
    needs_template_mapping = any(term in text for term in ("powerpoint", "pptx", "template", "slide"))
    has_template_mapping = any("template mapping" in component["service"].lower() for component in components)
    if needs_template_mapping and not has_template_mapping:
        components.append(
            {
                "id": "ARC-TEMPLATE-MAPPING",
                "service": "Template mapping layer",
                "classification": "required_for_poc",
                "purpose": "Populate predefined PowerPoint placeholders without changing template branding, layout, fonts, or colors.",
                "requirement_ids": ["functional_requirements"],
                "alternatives": ["Per-template JSON configuration"],
            }
        )
    return components


def _normalize_architecture(
    architecture: Dict[str, Any],
    extracted: Dict[str, List[str]],
    sources: List[Dict[str, str]],
    research_mode: str,
    private_references: Optional[List[Dict[str, str]]] = None,
) -> Dict[str, Any]:
    research_requested = research_mode != "off"
    usable_sources = _normalize_sources(sources, _normalize_research_mode(research_mode))
    research_status = "completed" if usable_sources else ("blocked" if research_requested else "not_requested")
    components = _ensure_architecture_coverage(
        _normalize_components(architecture.get("components"), extracted), extracted
    )
    selected_patterns = [pattern for pattern in architecture.get("selected_patterns", []) if pattern in ARCHITECTURE_PATTERN_NAMES]
    for pattern in _infer_architecture_patterns(extracted):
        if pattern not in selected_patterns:
            selected_patterns.append(pattern)

    def as_list(key: str) -> List[Any]:
        value = architecture.get(key, [])
        return value if isinstance(value, list) else [value] if value else []

    return {
        "research_status": research_status,
        "architecture_status": "draft" if research_status != "blocked" else "draft_pending_research",
        "architecture_summary": str(architecture.get("architecture_summary") or architecture.get("executive_recommendation") or "Architecture draft generated from confirmed requirements."),
        "selected_patterns": selected_patterns,
        "components": components,
        "security_design": as_list("security_design") or as_list("nfr_findings"),
        "scalability_design": as_list("scalability_design"),
        "monitoring_design": as_list("monitoring_design"),
        "cost_estimate": _normalize_cost_estimate(architecture.get("cost_estimate")),
        "architecture_decisions": as_list("architecture_decisions") or as_list("service_recommendations"),
        "decision_questions": as_list("decision_questions"),
        "risks": as_list("risks"),
        "research_sources": usable_sources,
        "private_reference_documents": [
            str(reference.get("document_name") or "reference-document")
            for reference in (private_references or [])
            if isinstance(reference, dict)
        ],
        "validation_notes": [
            "Architecture is a draft until a solution architect approves it.",
            "Model token usage must be reported separately from platform costs.",
        ] + (["Official web research was requested but no usable official source was returned."] if research_status == "blocked" else []),
    }


def research_architecture(
    extracted: Dict[str, List[str]],
    clarifications: Optional[Dict[str, str]] = None,
    research_mode: str = "azure_official_only",
    project_context: Optional[Dict[str, Any]] = None,
    private_references: Optional[List[Dict[str, str]]] = None,
    agent_requirements: Optional[Dict[str, Any]] = None,
) -> Dict[str, Any]:
    """Create a validated architecture draft with optional safe public research."""
    normalized_mode = _normalize_research_mode(research_mode)
    policy = _research_policy(normalized_mode)
    if published_agent_configured("architect"):
        try:
            raw_architecture = invoke_published_agent(
                "architect",
                {
                    "confirmed_requirements": agent_requirements or _requirements_with_ids(extracted),
                    "confirmed_hitl_answers": clarifications or {},
                    "project_context": project_context or {},
                    "research_policy": policy,
                },
            )
            architecture_payload = raw_architecture.get("architecture", raw_architecture)
            if isinstance(architecture_payload, dict) and not architecture_payload.get("components"):
                architecture_payload = {
                    **architecture_payload,
                    "components": architecture_payload.get("azure_services", []),
                }
            architecture = _normalize_architecture(
                architecture_payload,
                extracted,
                raw_architecture.get("sources", []),
                normalized_mode,
                private_references,
            )
            return {
                "mode": "foundry_published_architect_agent",
                "research_policy": policy,
                "architecture": architecture,
                "sources": architecture["research_sources"],
                "agent_output": raw_architecture,
            }
        except Exception as exc:
            if not azure_openai_configured():
                return {"mode": "failed", "error": str(exc), "architecture": {}, "sources": [], "research_policy": policy}

    if not azure_openai_configured():
        return {
            "mode": "unavailable",
            "error": "Configure the Published Architect Agent or Azure OpenAI Responses credentials before running architecture research.",
            "architecture": {},
            "sources": [],
            "research_policy": policy,
        }

    safe_requirements = {
        field: values for field, values in extracted.items() if field in REQUIREMENT_FIELDS
    }
    safe_web_queries = policy["safe_web_queries"]
    private_reference_summaries = [
        {
            "document_name": str(reference.get("document_name") or "reference-document"),
            "excerpt": str(reference.get("excerpt") or "")[:12000],
        }
        for reference in (private_references or [])
        if isinstance(reference, dict) and reference.get("excerpt")
    ][:3]
    prompt = f"""
You are an enterprise Azure solution architect for AI RFP proposals. Return valid JSON only.

Confirmed requirements JSON:
{json.dumps(safe_requirements, indent=2)}

Clarification answers JSON:
{json.dumps(clarifications or {}, indent=2)}

Project context JSON:
{json.dumps(project_context or {}, indent=2)}

Private reference excerpts JSON:
{json.dumps(private_reference_summaries, indent=2)}

Research mode: {normalized_mode}
Research policy JSON:
{json.dumps(policy, indent=2)}

Safe web queries, if research mode is not off: {json.dumps(safe_web_queries)}

RESEARCH POLICY
- When research mode is azure_official_only or azure_github_papers, invoke Web Search before writing the architecture.
- Search only the safe web queries above. Never put customer requirement text, customer names,
  private URLs, credentials, or internal identifiers into a Web Search query.
- Private reference excerpts are trusted project context. Never include their content in a Web Search query.
- For Azure claims, use only Microsoft first-party sources.
- For GitHub implementation evidence, use only github.com/Azure or github.com/microsoft repositories.
- For research-paper evidence, use only arxiv.org, aclanthology.org, openreview.net, or paperswithcode.com.
- Do not invent current capabilities, regions, limits, prices, SLAs, paper findings, repository behavior, or citations.
- If Web Search is unavailable, return a draft with research_status "blocked" and use "TBD"
  instead of unverified current facts or prices.

ARCHITECTURE RULES
- selected_patterns must use only: {json.dumps(ARCHITECTURE_PATTERN_NAMES)}.
- Include a pattern only when it is actually used.
- Map every active requirement category to a component or architecture decision.
- Use Azure AI Search only for grounded retrieval with citations and project metadata filters.
- Use native parsing first. Use Azure Document Intelligence only for scanned PDFs, tables, or complex layouts.
- Default to the lowest-cost POC component set. Mark API Management, Application Gateway,
  Front Door, AKS, VNet, Private Endpoint, and Private DNS as optional unless a requirement mandates them.
- Use Microsoft Entra ID. Enforce authorization in the backend and retrieval filters, never only in the frontend.
- If the requirements include a response-time target, preserve the exact target in scalability and monitoring design.
- If a PowerPoint template is required, include a template mapping layer that updates predefined placeholders only.
- Return Development, Test, and UAT/Staging cost structures. Use TBD unless a verified price source is available.
- Do not ask a decision question already answered in the clarification answers.

Return JSON only in this shape:
{{
  "research_status": "completed|blocked|not_requested",
  "architecture_summary": "string",
  "selected_patterns": ["RAG with metadata filtering"],
  "components": [{{
    "id": "ARC-001",
    "service": "Azure service or deterministic component name",
    "classification": "required_for_poc|optional_for_poc|production_only",
    "purpose": "string",
    "requirement_ids": ["requirement category or confirmed requirement ID"],
    "alternatives": ["string"]
  }}],
  "security_design": ["string"],
  "scalability_design": ["string"],
  "monitoring_design": ["string"],
  "cost_estimate": {{
    "development": [{{"service": "string", "monthly_estimate_usd": "TBD", "notes": "string"}}],
    "test": [{{"service": "string", "monthly_estimate_usd": "TBD", "notes": "string"}}],
    "uat_or_staging": [{{"service": "string", "monthly_estimate_usd": "TBD", "notes": "string"}}]
  }},
  "architecture_decisions": [{{"decision_id": "DEC-001", "decision": "string", "justification": "string", "requirement_ids": ["string"]}}],
  "decision_questions": [{{"id": "Q-001", "question": "string", "impact": "string", "requirement_ids": ["string"]}}],
  "risks": [{{"id": "R-001", "title": "string", "impact": "string", "mitigation": "string"}}]
}}

Keep the answer concise. Return no more than three decision questions.
"""
    try:
        research, citations = create_json_response(prompt, use_web_search=normalized_mode != "off")
        architecture = _normalize_architecture(
            research, safe_requirements, citations, normalized_mode, private_references
        )
        return {
            "mode": "azure_openai_web_search" if normalized_mode != "off" else "azure_openai",
            "research_policy": policy,
            "architecture": architecture,
            "sources": architecture["research_sources"],
        }
    except Exception as exc:
        return {"mode": "failed", "error": str(exc), "architecture": {}, "sources": [], "research_policy": policy}


def analyze_document(ingestion: Dict[str, Any], clarifications: Optional[Dict[str, str]] = None) -> Dict[str, Any]:
    extraction = extract_requirements_from_chunks(ingestion.get("chunks", []))
    extracted = extraction.get("extracted", {})
    questions = extraction.get("clarification_questions") or build_clarification_questions(extracted)
    return {
        "extraction_mode": extraction.get("mode", "llm"),
        "extracted": extracted,
        "traceability": extraction.get("traceability", {}),
        "agent_requirements": extraction.get("agent_requirements"),
        "ingestion": {
            "document_id": ingestion.get("document_id"),
            "filename": ingestion.get("filename"),
            "source_type": ingestion.get("source_type"),
            "stats": ingestion.get("stats", {}),
            "chunks": [
                {
                    "chunk_id": chunk.get("chunk_id"),
                    "section": chunk.get("section"),
                    "heading_path": chunk.get("heading_path"),
                    "page_start": chunk.get("page_start"),
                    "page_end": chunk.get("page_end"),
                    "estimated_tokens": chunk.get("estimated_tokens"),
                }
                for chunk in ingestion.get("chunks", [])
            ],
        },
        "batch_count": extraction.get("batch_count", 0),
        "provider_errors": extraction.get("provider_errors", []),
        "quality_guide": extraction.get("quality_guide", FIELD_QUALITY_GUIDE),
        "quality_notes": extraction.get("quality_notes", {}),
        "clarification_questions": questions,
    }


def read_uploaded_text(file_storage) -> str:
    filename = (file_storage.filename or "").lower()
    raw = file_storage.read()
    if filename.endswith(".txt") or filename.endswith(".md"):
        return raw.decode("utf-8", errors="ignore")
    if filename.endswith(".pdf"):
        try:
            from pypdf import PdfReader
            from io import BytesIO

            reader = PdfReader(BytesIO(raw))
            text = "\n".join(page.extract_text() or "" for page in reader.pages).strip()
            if text:
                return text
            return _read_pdf_with_document_intelligence(raw)
        except Exception as exc:
            try:
                return _read_pdf_with_document_intelligence(raw)
            except Exception:
                raise ValueError(f"PDF parsing failed. Use OCR or a text-based PDF. {exc}") from exc
    if filename.endswith(".docx"):
        try:
            from docx import Document
            from io import BytesIO

            doc = Document(BytesIO(raw))
            return "\n".join(paragraph.text for paragraph in doc.paragraphs)
        except Exception as exc:
            raise ValueError(f"DOCX parsing failed. {exc}")
    raise ValueError("Supported upload formats for the lightweight demo are .txt, .md, .pdf, and .docx.")


def _read_pdf_with_document_intelligence(raw: bytes) -> str:
    """Use Azure OCR/layout only after native PDF text extraction is unavailable."""
    endpoint = os.getenv("AZURE_DOCUMENT_INTELLIGENCE_ENDPOINT") or os.getenv("FORM_RECOGNIZER_ENDPOINT")
    key = os.getenv("AZURE_DOCUMENT_INTELLIGENCE_KEY") or os.getenv("FORM_RECOGNIZER_KEY")
    if not endpoint or not key:
        raise ValueError("The PDF has no extractable text and Document Intelligence is not configured.")

    from azure.ai.documentintelligence import DocumentIntelligenceClient
    from azure.core.credentials import AzureKeyCredential

    client = DocumentIntelligenceClient(endpoint, AzureKeyCredential(key))
    result = client.begin_analyze_document("prebuilt-layout", raw, content_type="application/pdf").result()
    text = (getattr(result, "content", "") or "").strip()
    if not text:
        raise ValueError("Document Intelligence returned no readable document content.")
    return text


def proposal_output_dir() -> Path:
    output_dir = Path(__file__).resolve().parent.parent / "generated"
    output_dir.mkdir(exist_ok=True)
    return output_dir
