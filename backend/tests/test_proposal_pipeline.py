import io
import sys
import time
import unittest
from pathlib import Path
from unittest.mock import patch

from docx import Document
from werkzeug.datastructures import FileStorage

BACKEND = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(BACKEND))

from proposal_agent import (
    analyze_document,
    extract_requirements_from_chunks,
    research_architecture,
    _normalize_architecture,
    _canonical_agent_requirement_records,
    _requirement_statements,
    _sanitize_extracted_requirements,
    fallback_generate_proposal,
    _normalize_sources,
)
from proposal_ingestion import CHUNK_MAX_TOKENS, choose_mode, estimate_tokens, ingest_text, ingest_upload
from pptx_tools import get_template_manifest, validate_slide_plan
import hitl_graph


class ProposalPipelineTests(unittest.TestCase):
    def test_docx_tables_are_preserved_in_ingestion(self):
        document = Document()
        document.add_heading("Security Requirements", level=1)
        document.add_paragraph("Employees must authenticate with single sign-on.")
        table = document.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "Retention"
        table.cell(0, 1).text = "90 days"
        table.cell(1, 0).text = "Audit logs"
        table.cell(1, 1).text = "12 months"
        raw = io.BytesIO()
        document.save(raw)
        raw.seek(0)

        ingestion = ingest_upload(FileStorage(raw, filename="requirements.docx"))

        self.assertEqual(ingestion["source_type"], "docx")
        self.assertGreaterEqual(ingestion["stats"]["chunks"], 1)
        joined = "\n".join(chunk["text"] for chunk in ingestion["chunks"])
        self.assertIn("Retention | 90 days", joined)
        self.assertIn("Audit logs | 12 months", joined)

    def test_image_only_pdf_uses_document_intelligence_ocr(self):
        from pypdf import PdfWriter

        raw = io.BytesIO()
        writer = PdfWriter()
        writer.add_blank_page(width=612, height=792)
        writer.write(raw)
        raw.seek(0)

        with patch(
            "proposal_ingestion._read_pdf_with_document_intelligence",
            return_value={1: "1 Scope\n\nThe platform must extract text from scanned RFP PDFs.\n[Table]\nField | Value\nRetention | 90 days"},
        ) as document_intelligence:
            ingestion = ingest_upload(FileStorage(raw, filename="scanned_rfp_ocr_case.pdf"))

        document_intelligence.assert_called_once()
        self.assertEqual(ingestion["source_type"], "pdf")
        self.assertEqual(ingestion["stats"]["parser_pages"]["document_intelligence_ocr"], 1)
        self.assertEqual(ingestion["stats"]["document_intelligence_pages"], 1)
        self.assertIn("scanned RFP PDFs", ingestion["chunks"][0]["text"])
        self.assertIn("Retention | 90 days", ingestion["chunks"][0]["text"])

    def test_table_heavy_pdf_uses_document_intelligence_layout(self):
        class FakeContentStream:
            def get_data(self):
                return b"".join([b"0 0 10 10 re S "] * 35)

        class FakePdfPage:
            images = []

            def extract_text(self):
                lines = [
                    "Requirement ID  Category  Priority  Description"
                ]
                lines.extend(
                    f"REQ-{index:03d}  Functional  High  Preserve row {index} source traceability."
                    for index in range(1, 12)
                )
                return "\n".join(lines)

            def get_contents(self):
                return FakeContentStream()

        class FakePdfReader:
            def __init__(self, _stream):
                self.pages = [FakePdfPage()]

        with patch("pypdf.PdfReader", FakePdfReader), patch(
            "proposal_ingestion._read_pdf_with_document_intelligence",
            return_value={1: "Requirement ID | Category | Priority | Description\nREQ-001 | Functional | High | Preserve source traceability."},
        ) as document_intelligence:
            ingestion = ingest_upload(FileStorage(io.BytesIO(b"%PDF fake"), filename="table_heavy_layout_case.pdf"))

        document_intelligence.assert_called_once()
        self.assertEqual(ingestion["stats"]["parser_pages"]["document_intelligence_layout"], 1)
        self.assertEqual(ingestion["stats"]["page_quality"][0]["reason"], "layout_risk")
        self.assertIn("Requirement ID | Category | Priority", ingestion["chunks"][0]["text"])

    def test_hierarchical_chunking_preserves_sections_and_pages(self):
        text = "\f".join(
            [
                "1 Scope\n\n" + ("The solution must accept customer requirement documents. " * 1600),
                "2 Security Requirements\n\n" + ("Employees must use single sign-on and authorized project access. " * 1600),
                "3 Performance\n\n" + ("The system must process a 150-page document within 12 minutes. " * 1600),
            ]
        )
        ingestion = ingest_text(text, "long-rfp.txt")
        chunks = ingestion["chunks"]

        self.assertEqual(ingestion["stats"]["extraction_mode"], "chunked")
        self.assertGreaterEqual(len(chunks), 2)
        self.assertTrue(all(chunk["chunk_id"] and chunk["section"] for chunk in chunks))
        self.assertEqual({1, 2, 3}, {chunk["page_start"] for chunk in chunks})
        self.assertTrue(all(chunk["estimated_tokens"] <= CHUNK_MAX_TOKENS for chunk in chunks))

    def test_small_documents_are_sent_as_whole_document_chunk(self):
        ingestion = ingest_text(
            "1 Scope\n\nThe solution must accept RFP documents.\n\n2 Security\n\nUsers must use single sign-on.",
            "small-rfp.txt",
        )

        self.assertEqual(choose_mode(ingestion["stats"]["estimated_tokens"]), "whole_document")
        self.assertEqual(ingestion["stats"]["extraction_mode"], "whole_document")
        self.assertEqual(len(ingestion["chunks"]), 1)
        self.assertEqual(ingestion["chunks"][0]["section"], "Whole Document")
        self.assertIn("2 Security", ingestion["chunks"][0]["text"])

    def test_large_documents_pack_near_heading_boundaries(self):
        text = "\f".join(
            [
                "1 Executive Summary\n\n" + ("The platform shall preserve proposal quality. " * 1300),
                "2 Functional Requirements\n\n" + ("Proposal managers shall upload and review RFP documents. " * 1300),
                "3 Security Requirements\n\n" + ("Employees must authenticate with enterprise single sign-on. " * 1300),
                "4 Reporting Requirements\n\n" + ("The system shall report token usage and approval history. " * 1300),
            ]
        )
        ingestion = ingest_text(text, "packed-rfp.txt")
        chunks = ingestion["chunks"]

        self.assertEqual(ingestion["stats"]["extraction_mode"], "chunked")
        self.assertLess(len(chunks), 4)
        self.assertTrue(any("through" in chunk["section"] for chunk in chunks))
        self.assertTrue(all(chunk["estimated_tokens"] <= CHUNK_MAX_TOKENS for chunk in chunks))

    def test_requirement_normalizer_removes_cross_category_duplicates(self):
        raw = {
            "functional_requirements": [
                "The solution shall read documents from Microsoft SharePoint.",
                "Architecture generation must not begin until an architect approves requirements.",
            ],
            "non_functional_requirements": [
                "Expected peak concurrency is 20 users.",
                "Minimize monthly operating cost.",
            ],
            "constraints": [
                "Use Microsoft Azure as the primary cloud platform.",
                "Documents must be retained for 90 days.",
            ],
            "integrations": ["Read documents from Microsoft SharePoint."],
            "security_needs": [
                "Employees must authenticate using single sign-on.",
                "Customer documents are confidential and only authorized project members may access them.",
            ],
            "expected_users": ["Expected peak concurrency is 20 users."],
        }
        normalized = _sanitize_extracted_requirements(raw)

        self.assertEqual(len(normalized["integrations"]), 1)
        self.assertEqual(normalized["expected_users"], [])
        self.assertIn("Expected peak concurrency is 20 users.", normalized["non_functional_requirements"])
        self.assertIn("Minimize monthly operating cost.", normalized["constraints"])
        self.assertIn("Employees must authenticate using single sign-on.", normalized["security_needs"])

    def test_canonical_agent_records_follow_deduplicated_categories(self):
        raw = {
            "functional_requirements": [
                {"id": "REQ-014", "statement": "The solution shall read documents from Microsoft SharePoint."},
                {"id": "REQ-015", "statement": "The solution shall read documents from Microsoft SharePoint."},
            ],
            "integrations": [
                {"id": "REQ-027", "statement": "Read documents from Microsoft SharePoint."},
            ],
        }
        extracted = _sanitize_extracted_requirements(_requirement_statements(raw))
        canonical = _canonical_agent_requirement_records(raw, extracted)

        self.assertEqual(canonical["functional_requirements"], [])
        self.assertEqual(len(canonical["integrations"]), 1)
        self.assertEqual(canonical["integrations"][0]["id"], "REQ-014")

    def test_analysis_returns_traceable_requirement_records(self):
        ingestion = ingest_text(
            "1 Security\n\nEmployees must authenticate using single sign-on. "
            "Customer documents are confidential and only authorized project members may access them.",
            "security.txt",
        )
        agent_result = {
            "requirements": {
                "security_needs": [
                    {
                        "id": "REQ-001",
                        "statement": "Employees must authenticate using single sign-on.",
                        "source_chunk_ids": [ingestion["chunks"][0]["chunk_id"]],
                        "confidence": "high",
                    }
                ]
            }
        }
        with patch("proposal_agent.published_agent_configured", return_value=True), patch(
            "proposal_agent.invoke_published_agent", return_value=agent_result
        ):
            analysis = analyze_document(ingestion)

        records = analysis["traceability"]["security_needs"]
        self.assertGreaterEqual(len(records), 1)
        self.assertTrue(records[0]["source_chunk_ids"])
        self.assertIn(records[0]["confidence"], {"high", "medium"})

    def test_requirement_agent_failure_never_uses_legacy_model(self):
        ingestion = ingest_text("Users must use single sign-on.", "security.txt")
        with patch("proposal_agent.published_agent_configured", return_value=False):
            analysis = analyze_document(ingestion)

        self.assertEqual(analysis["extraction_mode"], "failed")
        self.assertFalse(any(analysis["extracted"].values()))

    def test_published_requirement_agent_contract_preserves_ids_and_hitl(self):
        ingestion = ingest_text("Users must use single sign-on.", "security.txt")
        agent_result = {
            "requirements": {
                "security_needs": [
                    {
                        "id": "REQ-009",
                        "statement": "Users must use single sign-on.",
                        "source_chunk_ids": ["input-000"],
                        "confidence": "high",
                    }
                ]
            },
            "completion_form": [
                {
                    "id": "Q-001",
                    "requirement_id": "REQ-009",
                    "field": "identity_provider",
                    "question": "Which identity provider is required?",
                    "impact": "high",
                    "answer_type": "single_select",
                    "options": ["Microsoft Entra ID", "Okta"],
                }
            ],
        }
        with patch("proposal_agent.published_agent_configured", return_value=True), patch(
            "proposal_agent.invoke_published_agent", return_value=agent_result
        ):
            analysis = analyze_document(ingestion)

        self.assertEqual(analysis["extraction_mode"], "foundry_published_requirement_agent")
        self.assertEqual(analysis["agent_requirements"]["security_needs"][0]["id"], "REQ-009")
        self.assertEqual(analysis["traceability"]["security_needs"][0]["source_chunk_ids"], ["input-000"])
        self.assertEqual(analysis["clarification_questions"][0]["field"], "identity_provider")
        self.assertEqual(analysis["clarification_questions"][0]["options"], ["Microsoft Entra ID", "Okta"])

    def test_chunked_requirement_extraction_calls_agent_per_chunk_then_global_merges(self):
        chunks = [
            {
                "chunk_id": "doc-s01-c001",
                "section": "1 Security",
                "heading_path": ["1 Security"],
                "page_start": 1,
                "page_end": 5,
                "estimated_tokens": 10_000,
                "text": "Uploaded customer source documents must be deleted after 30 days.",
            },
            {
                "chunk_id": "doc-s02-c001",
                "section": "2 Retention",
                "heading_path": ["2 Retention"],
                "page_start": 6,
                "page_end": 9,
                "estimated_tokens": 10_000,
                "text": "Uploaded customer source documents must be retained for at least 90 days.",
            },
        ]

        def fake_invoke(agent_kind, payload):
            chunk = payload["chunks"][0]
            if chunk["chunk_id"] == "doc-s01-c001":
                return {
                    "requirements": {
                        "constraints": [
                            {
                                "id": "REQ-001",
                                "statement": "Uploaded customer source documents must be deleted after 30 days.",
                                "source_chunk_ids": ["doc-s01-c001"],
                                "confidence": "high",
                            }
                        ]
                    }
                }
            return {
                "requirements": {
                    "constraints": [
                        {
                            "id": "REQ-002",
                            "statement": "Uploaded customer source documents must be retained for at least 90 days.",
                            "source_chunk_ids": ["doc-s02-c001"],
                            "confidence": "high",
                        }
                    ]
                }
            }

        with patch("proposal_agent.published_agent_configured", return_value=True), patch(
            "proposal_agent.invoke_published_agent", side_effect=fake_invoke
        ) as invoke_mock:
            result = extract_requirements_from_chunks(chunks)

        self.assertEqual(invoke_mock.call_count, 2)
        self.assertEqual(result["batch_count"], 2)
        self.assertEqual(len(result["extracted"]["constraints"]), 2)
        self.assertEqual(result["clarification_questions"][0]["field"], "retention_policy")
        self.assertEqual(result["traceability"]["constraints"][0]["source_chunk_ids"], ["doc-s01-c001"])

    def test_published_architect_agent_contract_is_used(self):
        extracted = {field: [] for field in [
            "business_goals", "functional_requirements", "non_functional_requirements", "constraints",
            "assumptions", "integrations", "security_needs", "expected_users",
        ]}
        extracted["security_needs"] = ["Require single sign-on."]
        agent_requirements = {
            "security_needs": [{"id": "REQ-009", "statement": "Require single sign-on."}]
        }
        agent_result = {
            "architecture_summary": "Use Entra ID for enterprise authentication.",
            "selected_patterns": ["RAG with metadata filtering"],
            "azure_services": [
                {
                    "service": "Microsoft Entra ID",
                    "purpose": "Authenticate employees and enforce roles.",
                    "requirement_ids": ["REQ-009"],
                }
            ],
        }
        with patch("proposal_agent.published_agent_configured", return_value=True), patch(
            "proposal_agent.invoke_published_agent", return_value=agent_result
        ):
            result = research_architecture(extracted, agent_requirements=agent_requirements, research_mode="off")

        self.assertEqual(result["mode"], "foundry_published_architect_agent")
        self.assertIn("Microsoft Entra ID", [item["service"] for item in result["architecture"]["components"]])

    def test_architect_research_policy_allows_github_and_papers(self):
        extracted = {field: [] for field in [
            "business_goals", "functional_requirements", "non_functional_requirements", "constraints",
            "assumptions", "integrations", "security_needs", "expected_users",
        ]}
        extracted["functional_requirements"] = ["Extract requirements with source traceability."]
        agent_result = {
            "architecture_summary": "Use Azure services with evidence from docs, samples, and papers.",
            "selected_patterns": ["Workflow automation"],
            "components": [
                {
                    "service": "Azure AI Document Intelligence",
                    "purpose": "Extract OCR and layout from scanned PDFs.",
                    "requirement_ids": ["functional_requirements"],
                }
            ],
            "sources": [
                {"title": "Azure sample", "url": "https://github.com/Azure/azure-sdk-for-python"},
                {"title": "RAG paper", "url": "https://arxiv.org/abs/2005.11401"},
                {"title": "Untrusted blog", "url": "https://example.com/rag"},
            ],
        }
        with patch("proposal_agent.published_agent_configured", return_value=True), patch(
            "proposal_agent.invoke_published_agent", return_value=agent_result
        ) as invoke_mock:
            result = research_architecture(extracted, research_mode="azure_github_papers")

        policy = invoke_mock.call_args.args[1]["research_policy"]
        self.assertEqual(policy["mode"], "azure_github_papers")
        self.assertIn("research_paper", policy["allowed_source_types"])
        self.assertTrue(any("site:github.com/Azure" in query for query in policy["safe_web_queries"]))
        source_urls = [source["url"] for source in result["sources"]]
        self.assertIn("https://github.com/Azure/azure-sdk-for-python", source_urls)
        self.assertIn("https://arxiv.org/abs/2005.11401", source_urls)
        self.assertNotIn("https://example.com/rag", source_urls)

    def test_azure_official_research_filters_non_official_sources(self):
        sources = _normalize_sources(
            [
                "https://learn.microsoft.com/azure/search/search-security-overview",
                "https://github.com/microsoft/sample",
                "https://arxiv.org/abs/2005.11401",
                "https://example.com/article",
            ],
            "azure_official_only",
        )

        urls = [source["url"] for source in sources]
        self.assertIn("https://learn.microsoft.com/azure/search/search-security-overview", urls)
        self.assertIn("https://github.com/microsoft/sample", urls)
        self.assertNotIn("https://arxiv.org/abs/2005.11401", urls)
        self.assertNotIn("https://example.com/article", urls)

    def test_architecture_normalizer_enforces_contract(self):
        extracted = {
            "business_goals": ["Reduce proposal preparation time by 40 percent."],
            "functional_requirements": ["Export a PowerPoint template."],
            "non_functional_requirements": ["P95 response time below four seconds."],
            "constraints": ["Use Azure and minimize cost."],
            "assumptions": [],
            "integrations": ["Read documents from SharePoint."],
            "security_needs": ["Use SSO and project-scoped access."],
            "expected_users": ["Support 20 users."],
        }
        architecture = _normalize_architecture(
            {
                "architecture_summary": "Draft",
                "selected_patterns": ["RAG with metadata filtering"],
                "components": [
                    {
                        "service": "Azure Cognitive Search",
                        "purpose": "Grounded document retrieval.",
                        "requirement_ids": ["functional_requirements"],
                    },
                    {
                        "service": "Azure Functions (Consumption)",
                        "purpose": "Run the authenticated API and approval workflow.",
                        "requirement_ids": ["functional_requirements"],
                    },
                    {
                        "service": "Azure AD",
                        "purpose": "Authenticate enterprise users.",
                        "requirement_ids": ["security_needs"],
                    },
                ],
            },
            extracted,
            [{"title": "Azure AI Search", "url": "https://learn.microsoft.com/azure/search/"}],
            "azure_official_only",
        )

        self.assertEqual(architecture["research_status"], "completed")
        self.assertIn("Workflow automation", architecture["selected_patterns"])
        self.assertEqual(architecture["components"][0]["service"], "Azure AI Search")
        services = [item["service"] for item in architecture["components"]]
        self.assertIn("Microsoft Entra ID", services)
        self.assertNotIn("Azure Container Apps", services)
        self.assertEqual(set(architecture["cost_estimate"]), {"development", "test", "uat_or_staging"})
        self.assertTrue(any("template mapping" in item["service"].lower() for item in architecture["components"]))

    def test_150_page_text_ingestion_speed(self):
        page = "1 Requirement Section\n\n" + ("The platform must retain source metadata for traceability. " * 45)
        text = "\f".join(page for _ in range(150))

        started = time.perf_counter()
        ingestion = ingest_text(text, "150-pages.txt")
        elapsed = time.perf_counter() - started

        self.assertEqual(ingestion["stats"]["pages"], 150)
        self.assertGreater(ingestion["stats"]["estimated_tokens"], 40_000)
        self.assertEqual(ingestion["stats"]["extraction_mode"], "chunked")
        self.assertTrue(all(chunk["estimated_tokens"] <= CHUNK_MAX_TOKENS for chunk in ingestion["chunks"]))
        self.assertLess(elapsed, 3.0, f"Ingestion took {elapsed:.2f}s")

    def test_slide_planner_uses_problem_to_value_storyline_without_duplicate_components(self):
        proposal = fallback_generate_proposal(
            {
                "business_goals": ["Reduce proposal preparation time by 40 percent."],
                "functional_requirements": ["Export a customer-ready PowerPoint proposal."],
                "non_functional_requirements": ["Process a 100-page document within 10 minutes."],
                "constraints": ["Use Azure as the primary cloud platform."],
                "assumptions": [],
                "integrations": [],
                "security_needs": ["Require single sign-on and project-scoped access."],
                "expected_users": ["Support proposal managers and solution architects."],
            },
            architecture_decisions={
                "components": [
                    {"id": "ARC-LLM", "service": "Microsoft Foundry / Azure OpenAI", "purpose": "Generate approved proposal content."},
                    {"id": "ARC-LLM-2", "service": "Microsoft Foundry / Azure OpenAI", "purpose": "Duplicate should be removed."},
                    {"id": "ARC-PPTX", "service": "PPTX Renderer", "purpose": "Populate an approved enterprise template."},
                ],
                "selected_patterns": ["Workflow automation"],
            },
        )

        plan = proposal["slide_plan"]
        self.assertEqual(plan[0]["slide_type"], "customer_context")
        self.assertEqual(plan[3]["slide_type"], "solution_overview")
        self.assertEqual(plan[5]["slide_type"], "high_level_architecture")
        self.assertEqual(plan[-1]["slide_type"], "benefits_roadmap")
        self.assertTrue(all(slide["key_message"] for slide in plan))
        architecture_slide = next(slide for slide in plan if slide["slide_type"] == "high_level_architecture")
        names = [card["display_name"] for card in architecture_slide["component_cards"]]
        self.assertEqual(names.count("Microsoft Foundry / Azure OpenAI"), 1)
        self.assertTrue(all(len(slide["bullets"]) <= 5 for slide in plan))

    def test_starter_template_manifest_validates_generated_slide_plan(self):
        manifest = get_template_manifest("starter")
        self.assertEqual(manifest["template_id"], "starter")

        valid = validate_slide_plan(
            "starter",
            [{"slide_id": "S-01", "layout_id": "context", "title": "Customer Context", "bullets": ["A concise point."]}],
        )
        invalid = validate_slide_plan(
            "starter",
            [{"slide_id": "S-02", "layout_id": "context", "title": "Too dense", "bullets": ["a", "b", "c", "d", "e", "f"]}],
        )

        self.assertTrue(valid["valid"])
        self.assertFalse(invalid["valid"])
        self.assertEqual(invalid["errors"][0]["code"], "TOO_MANY_BULLETS")

    def test_langgraph_hitl_collects_all_questions_then_generates_only_after_approval(self):
        extraction = {
            "extraction_mode": "foundry_published_requirement_agent",
            "extracted": {"security_needs": ["Require single sign-on."], "functional_requirements": []},
            "agent_requirements": {
                "security_needs": [{"id": "REQ-009", "statement": "Require single sign-on."}],
                "functional_requirements": [],
            },
            "clarification_questions": [
                {"id": "Q-001", "field": "identity_provider", "question": "Which identity provider?", "answer_type": "single_select", "options": ["Microsoft Entra ID", "Okta"]},
                {"id": "Q-002", "field": "peak_concurrency", "question": "What peak concurrent user count?", "answer_type": "number"},
            ],
            "traceability": {"security_needs": []},
        }
        architect_result = {
            "mode": "foundry_published_architect_agent",
            "architecture": {
                "architecture_summary": "Use Microsoft Entra ID and Azure AI Search.",
                "selected_patterns": ["RAG with metadata filtering"],
                "cost_estimate": {"development": [{"service": "Azure AI Search", "monthly_estimate_usd": "0"}]},
            },
            "sources": [],
        }
        proposal = {"proposal": {"1_executive_summary": {"text": "Ready for export."}}}

        with patch("hitl_graph.analyze_document", return_value=extraction), patch(
            "hitl_graph.research_architecture", return_value=architect_result
        ) as research_mock, patch("hitl_graph.generate_proposal", return_value=proposal) as proposal_mock:
            started = hitl_graph.hitl_start("Employees need secure document Q&A.")
            self.assertEqual(started["state"]["status"], "awaiting_requirement_review")
            self.assertEqual(started["interrupt_payload"]["type"], "requirement_review")
            self.assertEqual(len(started["interrupt_payload"]["questions"]), 2)

            reviewed = hitl_graph.hitl_resume(started["thread_id"], {
                "decision": "approved",
                "clarifications": {"identity_provider": "Microsoft Entra ID", "peak_concurrency": "25"},
                "research_mode": "azure_github_papers",
                "additional_instruction": "Use GitHub samples and public papers for architecture evidence.",
            })
            self.assertEqual(reviewed["state"]["status"], "awaiting_architecture_approval")
            self.assertEqual(reviewed["interrupt_payload"]["type"], "architecture_review")
            self.assertEqual(research_mock.call_args.kwargs["agent_requirements"]["security_needs"][0]["id"], "REQ-009")
            self.assertEqual(research_mock.call_args.kwargs["clarifications"]["peak_concurrency"], "25")
            self.assertEqual(research_mock.call_args.kwargs["research_mode"], "azure_github_papers")
            self.assertIn("GitHub samples", research_mock.call_args.kwargs["clarifications"]["hitl_review"])
            self.assertEqual(reviewed["state"]["research_mode"], "azure_github_papers")
            proposal_mock.assert_not_called()

            completed = hitl_graph.hitl_resume(started["thread_id"], {"decision": "approved"})
            self.assertEqual(completed["state"]["status"], "completed")
            self.assertEqual(completed["state"]["proposal"], proposal)
            proposal_mock.assert_called_once()

    def test_langgraph_upload_preserves_docx_ingestion_for_requirement_agent(self):
        document = Document()
        document.add_paragraph("Proposal managers must upload documents for review.")
        table = document.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "Target"
        table.cell(0, 1).text = "10 minutes"
        raw = io.BytesIO()
        document.save(raw)
        raw.seek(0)
        extraction = {
            "extraction_mode": "foundry_published_requirement_agent",
            "extracted": {"functional_requirements": ["Upload documents for review."]},
            "agent_requirements": {"functional_requirements": [{"id": "REQ-003", "statement": "Upload documents for review."}]},
            "clarification_questions": [],
            "traceability": {},
        }

        with patch(
            "hitl_graph.analyze_document",
            side_effect=lambda ingestion: {**extraction, "ingestion": ingestion},
        ) as analyze_mock:
            started = hitl_graph.hitl_start_from_upload(FileStorage(raw, filename="requirements.docx"))

        supplied_ingestion = analyze_mock.call_args.args[0]
        self.assertEqual(supplied_ingestion["source_type"], "docx")
        self.assertIn("Target | 10 minutes", supplied_ingestion["chunks"][0]["text"])
        self.assertEqual(started["state"]["extraction"]["ingestion"]["source_type"], "docx")


if __name__ == "__main__":
    unittest.main()
